from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[3]
PPT_DIR = Path(__file__).resolve().parent
SRC = PPT_DIR / "Presentacion_1ra_Borrador_v03.pptx"
OUT = PPT_DIR / "Presentacion_1ra_Borrador_v04.pptx"


def remove_shape(shape):
    element = shape._element
    element.getparent().remove(element)


def remove_indices(slide, indices):
    for index in sorted(indices, reverse=True):
        remove_shape(slide.shapes[index])


def add_picture_cover(slide, image_path, left, top, width, height):
    pic = slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)
    image_ratio = pic.image.size[0] / pic.image.size[1]
    frame_ratio = width / height
    if image_ratio > frame_ratio:
        visible = frame_ratio / image_ratio
        crop = (1.0 - visible) / 2.0
        pic.crop_left = crop
        pic.crop_right = crop
    else:
        visible = image_ratio / frame_ratio
        crop = (1.0 - visible) / 2.0
        pic.crop_top = crop
        pic.crop_bottom = crop
    return pic


def add_picture_contain(slide, image_path, left, top, width, height):
    from PIL import Image

    with Image.open(image_path) as image:
        image_ratio = image.width / image.height
    frame_ratio = width / height
    if image_ratio > frame_ratio:
        final_width = width
        final_height = width / image_ratio
        final_left = left
        final_top = top + (height - final_height) / 2
    else:
        final_height = height
        final_width = height * image_ratio
        final_left = left + (width - final_width) / 2
        final_top = top
    return slide.shapes.add_picture(
        str(image_path), final_left, final_top, width=final_width, height=final_height
    )


prs = Presentation(SRC)

# Slide 22: replace the paper placeholder with the actual Lin & Ashlock figure.
slide = prs.slides[21]
box = slide.shapes[6]
left, top, width, height = box.left, box.top, box.width, box.height
remove_indices(slide, [8, 7, 6])
pic = slide.shapes.add_picture(str(PPT_DIR / "lin2016-p2.png"), left, top, width=width, height=height)
# Keep only the lower part of page 83, where Fig. 1 compares MSOR and MMSW.
pic.crop_top = 0.735
pic.crop_bottom = 0.005
pic.crop_left = 0.07
pic.crop_right = 0.07

# Slide 43: real hammer/transducer image already stored in the presentation project.
slide = prs.slides[42]
box = slide.shapes[5]
left, top, width, height = box.left, box.top, box.width, box.height
remove_indices(slide, [7, 6, 5])
add_picture_contain(
    slide,
    ROOT / "docs/Primera Presentación/latex-15p-review/figuras/excitacion/martillo_086d20.png",
    left,
    top,
    width,
    height,
)

# Slide 44: actual field-control UI with HAMMER and GEO traces.
slide = prs.slides[43]
box = slide.shapes[20]
left, top, width, height = box.left, box.top, box.width, box.height
remove_indices(slide, [22, 21, 20])
add_picture_cover(
    slide,
    ROOT / "docs/Primera Presentación/latex-15p-review/figuras/interfaces/geophone_scope_captura_real_claro.png",
    left,
    top,
    width,
    height,
)

# Slide 45: real aligned HAMMER/GEO acquisition used in the analysis workflow.
slide = prs.slides[44]
box = slide.shapes[5]
left, top, width, height = box.left, box.top, box.width, box.height
remove_indices(slide, [7, 6, 5])
add_picture_contain(
    slide,
    ROOT / "docs/Primera Presentación/latex-reestructurado/figuras/captura_cruda_campo.png",
    left,
    top,
    width,
    height,
)

# Slide 50: close with the strongest real project result instead of a placeholder.
slide = prs.slides[49]
box = slide.shapes[5]
left, top, width, height = box.left, box.top, box.width, box.height
remove_indices(slide, [7, 6, 5])
add_picture_cover(
    slide,
    ROOT / "docs/Primera Presentación/latex-15p-review/figuras/interfaces/masw_campo.png",
    left,
    top,
    width,
    height,
)

prs.save(OUT)
print(OUT)
