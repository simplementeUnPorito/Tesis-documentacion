"""v29 larga: contenido de la reestructura (después de restructure.ps1).
Puente (10), golpe (11), primera aproximación (12), lineal (17), curva G(γ) nueva (18),
interfaces (21) y Snell con ángulo crítico (22) antes del downhole, superficie libre (32)."""
import sys, copy, re
import pptx
from pptx.util import Emu
from lxml import etree

NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
      'a14': 'http://schemas.microsoft.com/office/drawing/2010/main'}
A = '{%s}' % NS['a']; M = '{%s}' % NS['m']
deck, fig = sys.argv[1], sys.argv[2]
prs = pptx.Presentation(deck)
S = lambda n: prs.slides[n - 1]

def title_is(n, start):
    t = S(n).shapes[0].text_frame.text if S(n).shapes[0].has_text_frame else ""
    assert t.startswith(start), f"slide {n}: «{t[:60]}» no empieza con «{start}»"

def set_text_all(slide, old, new):
    """Reemplaza en todos los a:t del slide, incluidas las copias Choice y Fallback."""
    n = 0
    for t in slide._element.iter(A + 't'):
        if t.text and old in t.text: t.text = t.text.replace(old, new); n += 1
    assert n, f"no se encontró «{old}»"
    return n

def shape(slide, name):
    return next(sh for sh in slide.shapes if sh.name == name)

def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def notes_sub(slide, old, new):
    tf = slide.notes_slide.notes_text_frame
    assert old in tf.text, f"nota sin «{old[:40]}»"
    tf.text = tf.text.replace(old, new)

# Párrafo de viñeta con runs normales, en cursiva y con subíndice: [(texto, 'n'|'i'|'s'|'si')]
def bullet(template_p, segs):
    p = copy.deepcopy(template_p)
    rpr = copy.deepcopy(p.find(A + 'r').find(A + 'rPr'))
    for r in p.findall(A + 'r'): p.remove(r)
    for ac in p.findall('mc:AlternateContent', NS): p.remove(ac)
    end = p.find(A + 'endParaRPr')
    for text, kind in segs:
        r = etree.SubElement(p, A + 'r'); rp = copy.deepcopy(rpr); r.append(rp)
        if 'i' in kind: rp.set('i', '1')
        if 's' in kind: rp.set('baseline', '-25000')
        etree.SubElement(r, A + 't').text = text
        if end is not None: end.addprevious(r)
    return p

def fallback_as_text(sp_):
    """La Fallback guardada por PowerPoint es una imagen del texto: se cambia por texto visible."""
    sppr = sp_.find('p:spPr', NS); bf = sppr.find(A + 'blipFill')
    if bf is not None: bf.addprevious(etree.Element(A + 'noFill')); sppr.remove(bf)
    for rp in sp_.iter(A + 'rPr'):
        nf = rp.find(A + 'noFill')
        if nf is not None:
            sf = etree.Element(A + 'solidFill'); etree.SubElement(sf, A + 'srgbClr').set('val', '1F4E79')
            nf.addprevious(sf); rp.remove(nf)

def ac_shapes(slide, name):
    """Las dos copias (Choice y Fallback) de un shape envuelto en mc:AlternateContent."""
    for ac in slide._element.iterfind('.//p:spTree/mc:AlternateContent', NS):
        names = [e.get('name') for e in ac.iter('{%s}cNvPr' % NS['p'])]
        if name in names: return ac, ac.findall('.//p:sp', NS)
    raise KeyError(name)

# ---------- 5 y 10: mapa ----------
for n in (5, 10):
    title_is(n, "Hay dos formas de medir el terreno")
    set_text_all(S(n), "por su resistencia o por cómo transmite ondas",
                 "por su resistencia a la penetración o por cómo se propagan las ondas en él")

s = S(10)
shape(s, "Título").text_frame.paragraphs[0].runs[0].text = "Para medir la rigidez sin recurrir a correlaciones, se usan las ondas"
body = shape(s, "Cuerpo").text_frame.paragraphs
body[0].runs[0].text = "Penetración: entrega resistencia, y la rigidez sale de correlaciones"
body[1].runs[0].text = "Ondas: su velocidad depende de la rigidez del terreno"
for nm in ("Fila resistencia", "Celda penetración dentro", "Celda penetración superficie"):
    sh = shape(s, nm); sh.fill.solid(); sh.fill.fore_color.rgb = pptx.dml.color.RGBColor(0xF2, 0xF2, 0xF2)
for nm in ("Text 6", "Text 10", "Text 12"):
    for p in shape(s, nm).text_frame.paragraphs:
        for r in p.runs: r.font.color.rgb = pptx.dml.color.RGBColor(0xA6, 0xA6, 0xA6)
set_notes(s, "Las dos técnicas de penetración entregan resistencia, y la rigidez sale de correlaciones. "
             "Para medirla de forma directa hay que pasar a la otra fila del mapa: la propagación de ondas. "
             "La velocidad de una onda depende de la rigidez del medio por el que se propaga, y eso es lo que vamos a aprovechar. "
             "Primero veamos qué es una onda en el terreno.")

# ---------- 11: el golpe ----------
title_is(11, "Al excitar un terreno sólido")
shape(S(11), "Título").text_frame.paragraphs[0].runs[0].text = "Un golpe deforma el terreno por un instante, y esa deformación se propaga"
set_text_all(S(11), "Cada tipo de deformación viaja como una onda", "Cada tipo de deformación se propaga como una onda")

# ---------- 12: primera aproximación ----------
title_is(12, "Para analizarlo, simplificamos")
shape(S(12), "Título").text_frame.paragraphs[0].runs[0].text = "Primera aproximación: un medio homogéneo, isótropo y elástico lineal"
notes_sub(S(12), "Para poder analizar el problema, lo simplificamos hasta el absurdo.",
          "Como primera aproximación, simplificamos mucho el problema.")

# ---------- 17: tramo lineal ----------
title_is(17, "En suelos, la relación esfuerzo-deformación")
set_text_all(S(17), "Esfuerzo y deformación de corte son proporcionales", "En pequeña deformación, τ y γ son proporcionales")
notes_sub(S(17), "Fuera de ese rango, G baja con la deformación.",
          "Fuera de ese rango, G baja con la deformación, como muestra la slide siguiente.")

# ---------- 18: curva G(γ) (copia de la 17) ----------
s = S(18); title_is(18, "En suelos, la relación esfuerzo-deformación")
shape(s, "Título").text_frame.paragraphs[0].runs[0].text = "Al crecer la deformación, la rigidez del suelo cae desde un máximo"
old = shape(s, "Corte simple")
pic = s.shapes.add_picture(fig, old.left, old.top, old.width, old.height)
pic.name = "Curva G-gamma"; old._element.addprevious(pic._element); old._element.getparent().remove(old._element)

# ecuación: G(γ) = G_max / (1 + γ/γ_r)
ac, (ch_sp, fb_sp) = ac_shapes(s, "Ecuación E_tau")
mr0 = ch_sp.find('.//m:r', NS)
def mr(text, upright=False):
    r = copy.deepcopy(mr0); r.find(M + 't').text = text
    if upright:
        rp = etree.Element(M + 'rPr'); etree.SubElement(rp, M + 'sty').set(M + 'val', 'p'); r.insert(0, rp)
        r.find(A + 'rPr').attrib.pop('i', None)
    return r
def sub(base, s_, upright_sub=False):
    e = etree.Element(M + 'sSub')
    etree.SubElement(e, M + 'e').append(mr(base)); etree.SubElement(e, M + 'sub').append(mr(s_, upright_sub))
    return e
def frac(num, den, lin=False):
    f = etree.Element(M + 'f')
    if lin: etree.SubElement(etree.SubElement(f, M + 'fPr'), M + 'type').set(M + 'val', 'lin')
    nu = etree.SubElement(f, M + 'num'); de = etree.SubElement(f, M + 'den')
    for x in num: nu.append(x)
    for x in den: de.append(x)
    return f
om = ch_sp.find('.//m:oMath', NS)
for c in list(om): om.remove(c)
om.append(mr("G(γ)="))
om.append(frac([sub("G", "max", True)], [mr("1+"), frac([mr("γ")], [sub("γ", "r")], lin=True)]))
for sp_ in (ch_sp, fb_sp):
    for c in sp_.iter('{%s}cNvPr' % NS['p']): c.set('name', 'Ecuación E_hiperbolica')
fb_sp.find('.//a:t', NS).text = "G(γ) = Gmax / (1 + γ/γr)"
fallback_as_text(fb_sp)

# viñetas: se reemplaza el Cuerpo (con fórmulas inline) por un cuadro de texto simple
ac_b, (chb, fbb) = ac_shapes(s, "Cuerpo")
# plantilla: la copia Choice (la Fallback que guarda PowerPoint es una imagen del texto viejo)
new_sp = copy.deepcopy(chb); tx = new_sp.find('.//p:txBody', NS)
tpl = tx.find('a:p', NS)
for p in tx.findall('a:p', NS): tx.remove(p)
for segs in ([("γ", "i"), ("r", "si"), (": deformación de referencia, que depende del suelo", "n")],
             [("Si ", "n"), ("γ", "i"), (" ≪ ", "n"), ("γ", "i"), ("r", "si"), (", ", "n"), ("G", "i"), (" ≈ ", "n"),
              ("G", "i"), ("max", "s"), (": la relación es lineal", "n")],
             [("Al crecer ", "n"), ("γ", "i"), (", ", "n"), ("G", "i"), (" cae: ", "n"), ("G", "i"), ("max", "s"),
              (" es su valor máximo", "n")]):
    tx.append(bullet(tpl, segs))
ac_b.addprevious(new_sp); ac_b.getparent().remove(ac_b)
shape(s, "Cita").text_frame.paragraphs[0].runs[0].text = "Modelo hiperbólico: Hardin y Drnevich (1972). Deformación de referencia ilustrativa: 10⁻³."
set_notes(s, "¿Qué pasa fuera de ese rango? Un modelo clásico para ilustrarlo es el hiperbólico de Hardin y Drnevich: "
             "G es Gmax dividido por uno más gamma sobre la deformación de referencia. "
             "Para deformaciones mucho menores que la de referencia, G queda prácticamente en Gmax: es el tramo lineal. "
             "Al crecer la deformación, G cae, y cuando gamma iguala a la de referencia, G vale la mitad de Gmax. "
             "Por eso a la rigidez en pequeña deformación se la llama Gmax: es el máximo de la curva. "
             "Las ondas de prospección trabajan en la zona sombreada, a la izquierda. "
             "Nota para mí: la deformación de referencia de 10 a la menos 3 es ilustrativa y depende del suelo; "
             "con ese valor, en 10 a la menos 5 la rigidez conserva el 99 %.")

# ---------- 21: interfaces, ahora antes del downhole ----------
title_is(21, "Las ondas de cuerpo, al llegar a una interfaz")
notes_sub(S(21), "Hasta acá el medio era uniforme. Pero el terreno real tampoco es homogéneo:",
          "Hasta acá el medio era homogéneo, como en la primera aproximación. Pero el terreno real no lo es:")
notes_sub(S(21), "Ese es el caso general. Ahora el caso que más nos interesa: cuando el segundo medio es el aire.",
          "Para saber hacia dónde sale cada onda, usamos la ley de Snell.")

# ---------- 22: Snell, sin SV/SH (se definen después) y con ángulo crítico ----------
s = S(22); title_is(22, "Los ángulos de las ondas generadas")
set_text_all(s, "Vale para P, SV y SH: V es ", "Vale para P y S: V es ")
set_text_all(s, " para SV y SH", " para S")
ac_s, (chs, fbs) = ac_shapes(s, "Cuerpo")
for sp_ in (chs, fbs):
    tx = sp_.find('.//p:txBody', NS); tpl = tx.findall('a:p', NS)[0]
    tx.append(bullet(tpl, [("Ángulo crítico: con ", "n"), ("V", "i"), ("2", "s"), (" > ", "n"), ("V", "i"), ("1", "s"),
                           (" y sin ", "n"), ("θ", "i"), ("c", "s"), (" = ", "n"), ("V", "i"), ("1", "s"), ("/", "n"),
                           ("V", "i"), ("2", "s"), (", la onda transmitida viaja por la interfaz", "n")]))
set_notes(s, "La ley de Snell relaciona los ángulos de todas las ondas que se generan en una interfaz: "
             "el seno del ángulo dividido por la velocidad es constante, y ese valor se llama parámetro del rayo. "
             "Cada onda usa su propia velocidad: Vp para la P y Vs para la S. Como la S es más lenta que la P, sale con un ángulo menor. "
             "Si el segundo medio es más rápido, la onda transmitida se acerca a la interfaz a medida que crece el ángulo de incidencia. "
             "Hay un ángulo en el que sale rasante y viaja por la interfaz: es el ángulo crítico, y su seno es V1 sobre V2. "
             "Ese ángulo es la base del método de refracción, que vemos enseguida.")

# ---------- 23: downhole ----------
title_is(23, "Downhole")
notes_sub(S(23), "Con Vs ya definida, vemos los métodos que la miden de forma directa.",
          "Con Vs definida y el terreno en capas, vemos los métodos que miden Vs de forma directa.")

# ---------- 32: superficie libre ----------
s = S(32); title_is(32, "En el caso de una superficie libre")
shape(s, "Título").text_frame.paragraphs[0].runs[0].text = "La superficie libre es una interfaz especial: casi toda la energía vuelve al terreno"
notes_sub(s, "Tomemos el caso en que el segundo medio es el aire.",
          "Volvamos a las interfaces con un caso especial: el segundo medio es el aire.")

# ---------- números de slide que se corrieron (+2 desde la 10) ----------
for sl in prs.slides:
    for t in sl._element.iter(A + 't'):
        if t.text:
            t.text = t.text.replace("modelo de la slide 39", "modelo de la slide 41").replace("curva de la slide 41", "curva de la slide 43")
    if sl.has_notes_slide and "slide 54" in sl.notes_slide.notes_text_frame.text:
        notes_sub(sl, "slide 54", "slide 56")

# ---------- referencia nueva ----------
ref = "Hardin, B. O. y Drnevich, V. P. (1972). Shear modulus and damping in soils: design equations and curves. J. Soil Mech. Found. Div., ASCE, 98(SM7), 667–692."
done = False
for sl in prs.slides:
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        for p in sh.text_frame.paragraphs:
            if p.text.startswith("Hakimitoroghi, N., Raut"):
                newp = copy.deepcopy(p._p); p._p.addnext(newp)
                rs = newp.findall(A + 'r'); rs[0].find(A + 't').text = ref
                for r in rs[1:]: newp.remove(r)
                done = True; break
assert done, "no se encontró dónde insertar la referencia"

prs.save(deck)
print("ok")

