"""v35 larga: diagramas de control en 64 (balance de fuerza) y 68 (AFE), viñeta de ganancias con los K."""
import sys, os
import pptx
from pptx.util import Pt
from pptx_helpers_v30 import NS, ptext, top_shapes, plain_par, body_par

deck = sys.argv[1]
FIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figuras_es_v25")
prs = pptx.Presentation(deck)
S = lambda n: prs.slides[n - 1]
title = lambda n: ptext(top_shapes(S(n), "Título")[1][0].find('.//a:p', NS))

def swap(n, name, img, box=None):
    old = next(sh for sh in S(n).shapes if sh.name == name)
    x, y, w, h = box or (old.left, old.top, old.width, old.height)
    pic = S(n).shapes.add_picture(img, x, y)
    r = min(w / pic.width, h / pic.height); pic.width, pic.height = int(pic.width * r), int(pic.height * r)
    pic.left, pic.top = int(x + (w - pic.width) / 2), int(y + (h - pic.height) / 2); pic.name = name
    old._element.addprevious(pic._element); old._element.getparent().remove(old._element)

# 68: AFE con triángulos y K simbólicos
assert title(68).startswith("Implementación del AFE"), title(68)
swap(68, "Diagrama de bloques del AFE", os.path.join(FIG, "fig_afe_bloques_control_es.png"))
plain_par(body_par(S(68), "Cuerpo", "Ganancias")[1],
          [("Ganancias: ", "n"), ("K", "i"), ("dif", "s"), (" = 2, ", "n"), ("K", "i"), ("PGA", "s"), (" de 1 a 50, ", "n"),
           ("K", "i"), ("sum", "s"), (" = 3,6 y ", "n"), ("K", "i"), ("LP", "s"), (" = 5: hasta ×1800 en total", "n")])

# 64: balance de fuerza como lazo de control; ocupa el área que deja libre la columna de viñetas
assert title(64).startswith("Balance de fuerza"), title(64)
old = next(sh for sh in S(64).shapes if sh.name == "Lazo de balance de fuerza")
swap(64, "Lazo de balance de fuerza", os.path.join(FIG, "fig_balance_fuerza_es.png"),
     box=(old.left, old.top, old.width, old.height))

prs.save(deck); print("ok")

# 64: más ancho para el diagrama; la columna de viñetas se angosta y baja a 16 pt
def layout_64(prs):
    from pptx_helpers_v30 import unwrap, A
    s = prs.slides[63]
    pic = next(sh for sh in s.shapes if sh.name == "Lazo de balance de fuerza")
    tit = next(sh for sh in s.shapes if sh.name == "Título"); msg = next(sh for sh in s.shapes if sh.name == "Text 4")
    top, bottom = tit.top + tit.height + Pt(6), msg.top - Pt(8)
    w = Pt(420); r = w / pic.width; pic.width, pic.height = int(w), int(pic.height * r)
    pic.left = Pt(30); pic.top = top + (bottom - top - pic.height) // 2
    cue = unwrap(s, "Cuerpo")
    off = cue.find('.//a:xfrm/a:off', NS); ext = cue.find('.//a:xfrm/a:ext', NS)
    off.set('x', str(Pt(462))); off.set('y', str(top)); ext.set('cx', str(Pt(684) - Pt(462))); ext.set('cy', str(bottom - top))
    for rp in cue.iter(A + 'rPr'): rp.set('sz', '1600')
    for rp in cue.iter(A + 'endParaRPr'): rp.set('sz', '1600')
    from pptx_helpers_v30 import plain_par, body_par
    plain_par(body_par(s, "Cuerpo", "El lazo interactúa")[1], "Hay que resolver estabilidad, ruido y tolerancias entre unidades")
    for sp_ in cue.iter(A + 'spcPts'): sp_.set('val', '400')

if __name__ == "__main__":
    prs = pptx.Presentation(deck); layout_64(prs); prs.save(deck); print("64 ok")
