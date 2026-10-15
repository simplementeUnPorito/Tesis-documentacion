"""v30 larga, paso 3 (después de restructure_v30_larga.ps1): slides nuevas de capas (21) y sensores (54),
interfaces sin lo que pasó a capas (22), números de slide corridos y referencias repartidas en cinco slides."""
import sys, copy, re, unicodedata
import pptx
from pptx_helpers_v30 import NS, A, P, ptext, top_shapes, unwrap, plain_par, body_par, set_title, set_notes

deck, fig_capas = sys.argv[1], sys.argv[2]
prs = pptx.Presentation(deck)
S = lambda n: prs.slides[n - 1]
title = lambda n: ptext(top_shapes(S(n), "Título")[1][0].find('.//a:p', NS))

# ---------------- 21: capas (copia de interfaces) ----------------
s = S(21); assert title(21).startswith("Al llegar a una interfaz")
set_title(s, "Del medio homogéneo al modelo en capas")
old = next(sh for sh in s.shapes if sh.name == "Figura interfaz")
pic = s.shapes.add_picture(fig_capas, old.left, old.top, height=old.height)
pic.left = old.left + (old.width - pic.width) // 2 if pic.width < old.width else old.left
pic.name = "Figura capas"; old._element.addprevious(pic._element); old._element.getparent().remove(old._element)
sp, p_first = body_par(s, "Cuerpo", "Las ondas de cuerpo")
for start in ("Las ondas de cuerpo", "Parte de la energía"):
    _, p = body_par(s, "Cuerpo", start); p.getparent().remove(p)
_, p_last = body_par(s, "Cuerpo", "Un modelo en capas")
new = copy.deepcopy(body_par(s, "Cuerpo", "El terreno real")[1]); p_last.addnext(new)
plain_par(new, "Dentro de cada capa se mantiene la primera aproximación")
set_notes(s, "La primera aproximación suponía un medio homogéneo, pero el terreno real no lo es: varía de forma continua "
             "y tiene discontinuidades. Un paso más realista es dividirlo en capas. Dentro de cada capa se mantiene la "
             "primera aproximación, y en cada interfaz cambian la densidad y las velocidades de las ondas P y S. "
             "Así, el problema se reduce a entender qué pasa cuando una onda llega a una interfaz.")

# ---------------- 22: interfaces ----------------
s = S(22); assert title(22).startswith("Al llegar a una interfaz")
for start in ("El terreno real", "Un modelo en capas"):
    _, p = body_par(s, "Cuerpo", start); p.getparent().remove(p)
_, p = body_par(s, "Cuerpo", "Parte de la energía")
new = copy.deepcopy(p); p.addnext(new); plain_par(new, "Cuánto se refleja depende del contraste entre los dos medios")
set_notes(s, "Cuando una onda de cuerpo, P o S, llega a una interfaz, una parte de la energía se refleja y otra se "
             "transmite al segundo medio. Cuánto se refleja depende del contraste entre los dos medios. "
             "Para saber hacia dónde sale cada onda, se usa la ley de Snell.")

# ---------------- 54: transductores posibles (copia de la tabla de ESP-NOW) ----------------
s = S(54); assert title(54).startswith("Se eligió ESP-NOW")
set_title(s, "Transductores posibles para la componente vertical")
tbl = next(sh for sh in s.shapes if sh.has_table).table._tbl
trs = tbl.tr_lst
order = [trs[0], copy.deepcopy(trs[5]), trs[1], trs[2], trs[3]]  # la fila resaltada (ESP-NOW) pasa a ser la del SM-24
for tr in list(trs): tbl.remove(tr)
for tr in order: tbl.append(tr)
rows = [("Transductor", "Banda", "Precio (USD)", "Observación"),
        ("Geófono SM-24", "fn = 10 Hz", "69,95", "Pasivo. Atenúa por debajo de fn"),
        ("Geófono RTC 4,5 Hz", "fn = 4,5 Hz", "92–154", "Baja la resonancia, pero conserva la caída por debajo de ella"),
        ("Acelerómetro MEMS ADXL355", "Desde continua", "70,93", "Mide aceleración. Su ruido no se ensayó para esta campaña"),
        ("Sismómetro Güralp 3ESPC", "0,017–100 Hz", "Por cotización", "Banda amplia. Instrumento de 8,4 kg, no un elemento sensor")]
for tr, vals in zip(tbl.tr_lst, rows):
    for tc, v in zip(tr.tc_lst, vals):
        p = tc.find('.//a:p', NS); plain_par(p, v)
        for extra in tc.findall('.//a:p', NS)[1:]: extra.getparent().remove(extra)
txt = next(sh for sh in s.shapes if sh.name == "Text 1")
txt.text_frame.paragraphs[0].runs[0].text = "Se eligió el SM-24 por disponibilidad. Su caída por debajo de 10 Hz se compensa con electrónica."
for r in txt.text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
cita = next(sh for sh in s.shapes if sh.name == "Cita")
cita.text_frame.paragraphs[0].runs[0].text = ("SM-24: Sensor Nederland B.V. (2006). RTC: R. T. Clark (2026). ADXL355: Analog Devices (2026). "
                                              "3ESPC: Güralp Systems (s. f.). Precios del elemento sensor al 31/08/2026: "
                                              "SparkFun Electronics (2026), R. T. Clark (2026) y DigiKey (2026).")
set_notes(s, "Antes de describir el nodo hay que elegir el transductor que convierte el movimiento vertical del suelo en una "
             "señal eléctrica. Un geófono de 4,5 Hz baja la resonancia, pero conserva la caída por debajo de ella y cuesta "
             "más por canal. Un acelerómetro MEMS como el ADXL355 mide desde continua, pero no se ensayó si su ruido alcanza "
             "para estas señales. Un sismómetro de banda ancha como el 3ESPC cubre todo, pero es un instrumento de 8,4 kg "
             "que se cotiza aparte. Se eligió el SM-24 porque estaba disponible, y su límite en baja frecuencia se compensa "
             "con electrónica, que es el tema central del trabajo.")

# ---------------- números de slide que se corrieron ----------------
for sl in prs.slides:
    for t in sl._element.iter(A + 't'):
        if t.text:
            t.text = t.text.replace("modelo de la slide 41", "modelo de la slide 42").replace("curva de la slide 43", "curva de la slide 44")
    if sl.has_notes_slide:
        tf = sl.notes_slide.notes_text_frame
        if "slide 56" in tf.text: tf.text = tf.text.replace("slide 56", "slide 58")

# ---------------- referencias: tres nuevas, orden alfabético, cinco slides ----------------
NEW = ["Analog Devices (2026). ADXL355: Low-noise, low-drift, low-power 3-axis MEMS accelerometer. Página de producto y hoja de datos. analog.com.",
       "DigiKey (2026). ADXL355BEZ-RL7: pricing and availability. digikey.com.",
       "Güralp Systems (s. f.). 3ESPC portable broadband seismometer. Hoja de datos DAS-C3E-0001. guralp.com."]
boxes = [sh for sl in prs.slides for sh in sl.shapes if sh.name == "Referencias"]
assert len(boxes) == 5, len(boxes)
pars = [p._p for b in boxes for p in b.text_frame.paragraphs]
key = lambda e: unicodedata.normalize('NFKD', "".join(t.text or "" for t in e.iter(A + 't'))).encode('ascii', 'ignore').decode().lower()
for text in NEW:
    e = copy.deepcopy(pars[0]); rs = e.findall(A + 'r')
    rs[0].find(A + 't').text = text
    for r in rs[1:]: e.remove(r)
    pars.append(e)
pars.sort(key=key)
for b in boxes:
    tx = b.text_frame._txBody
    for p in tx.findall(A + 'p'): tx.remove(p)
total = sum(len(key(e)) + 60 for e in pars); acc = 0; k = 0
for e in pars:
    if acc > total * (k + 1) / 5 and k < 4: k += 1
    boxes[k].text_frame._txBody.append(e); acc += len(key(e)) + 60
for k, b in enumerate(boxes, 1):
    sl = next(sl for sl in prs.slides if b in list(sl.shapes))
    tsp = top_shapes(sl, "Título")[1][0]
    for t in tsp.iter(A + 't'):
        if t.text and "Referencias" in t.text: t.text = re.sub(r"\(\d de \d\)", f"({k} de 5)", t.text)

prs.save(deck)
print("ok", [len(b.text_frame.paragraphs) for b in boxes])
