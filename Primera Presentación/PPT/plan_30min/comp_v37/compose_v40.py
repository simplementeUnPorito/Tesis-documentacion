"""Compositor de la comprimida v37 (un tema por slide, coreografía de anim3).
Uso: python compose.py comp29.pptx choreo29.json larga37.pptx textmap.json spec.py salida.pptx anim3_dir
- Escenas «old»: objetos y rects de la comprimida v29 (choreo original): zoom, swap, persistencia.
- Escenas «new»: contenido de una slide de la larga v37, ubicado en el área de escena de la comprimida.
- Escenas «keys»: resumen en palabras clave de color (chips) al cierre de cada tema.
Las relaciones se reapuntan (mismo paquete) o se copian (imágenes desde la larga)."""
import sys, os, json, copy, re, io, importlib.util
import pptx
from lxml import etree
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, os.path.dirname(HERE))
from pptx_helpers_v30 import NS, P, A, M, MC, in_fallback, ptext, plain_par

R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
EMU = 914400
comp_path, choreo_path, larga_path, tmap_path, spec_path, out_path, anim_dir = sys.argv[1:8]
sys.path.insert(0, anim_dir); import anim3                           # coreografía original
anim3.rect_diff = lambda a, b: max(abs((a[i] or 0.0) - (b[i] or 0.0)) for i in range(4))   # None vale 0, como en la pasada final
spec = importlib.util.spec_from_file_location("spec", spec_path); SPEC = importlib.util.module_from_spec(spec); spec.loader.exec_module(SPEC)

comp = pptx.Presentation(comp_path); larga = pptx.Presentation(larga_path)
CH = json.load(open(choreo_path, encoding='utf-8'))
TMAP = json.load(open(tmap_path, encoding='utf-8'))
old_slides = list(comp.slides); L = list(larga.slides)

# ---------------------------------------------------------------- utilidades
def tag_of(el):
    for c in el.iter(P + 'cNvPr'):
        m = re.search(r'\[o(\d+)\]', c.get('name', ''))
        if m: return int(m.group(1))
    return None

def top_children(slide):
    tree = slide._element.find('.//p:spTree', NS)
    return [ch for ch in tree if ch.tag not in (P + 'nvGrpSpPr', P + 'grpSpPr')]

def tagged(slide):
    d = {}
    for el in top_children(slide):
        t = tag_of(el)
        if t is not None: d.setdefault(t, el)
    return d

def copy_el(el, src_part, dst_part, cross):
    new = copy.deepcopy(el)
    for node in new.iter():
        for attr in list(node.attrib):
            if not attr.startswith(R): continue
            rid = node.get(attr); rel = src_part.rels[rid]
            if rel.is_external:
                nid = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            elif not cross:
                nid = dst_part.relate_to(rel.target_part, rel.reltype)
            elif rel.reltype.endswith('/image'):
                _, nid = dst_part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
            else:
                raise RuntimeError(f"relación no copiable entre paquetes: {rel.reltype}")
            node.set(attr, nid)
    return new

def set_tag(el, tag):
    for c in el.iter(P + 'cNvPr'):
        nm = re.sub(r'\s*\[o\d+\]', '', c.get('name', ''))
        c.set('name', f"{nm} [o{tag}]".strip())

def xfrm_of(el):
    """xfrm de nivel superior (en AlternateContent, el de Choice y el de Fallback)."""
    if el.tag == MC + 'AlternateContent':
        return [x for sp in el.findall('.//p:sp', NS) for x in [sp.find('p:spPr/a:xfrm', NS)] if x is not None]
    if el.tag == P + 'graphicFrame': return [el.find('p:xfrm', NS)]
    if el.tag == P + 'grpSp': return [el.find('p:grpSpPr/a:xfrm', NS)]
    return [x for x in [el.find('p:spPr/a:xfrm', NS)] if x is not None]

def rect_in(el):
    x = xfrm_of(el)[0]; o = x.find('a:off', NS); e = x.find('a:ext', NS)
    return [int(o.get('x')) / EMU, int(o.get('y')) / EMU, int(e.get('cx')) / EMU, int(e.get('cy')) / EMU]

def set_rect(el, r):
    for x in xfrm_of(el):
        o = x.find('a:off', NS); e = x.find('a:ext', NS)
        o.set('x', str(int(r[0] * EMU))); o.set('y', str(int(r[1] * EMU)))
        e.set('cx', str(int(r[2] * EMU))); e.set('cy', str(int(r[3] * EMU)))

def scale_fonts(el, s):
    for node in el.iter(A + 'rPr', A + 'endParaRPr', A + 'defRPr'):
        if node.get('sz'): node.set('sz', str(max(800, int(round(int(node.get('sz')) * s / 100.0)) * 100)))
    for tr in el.iter(A + 'tr'):
        if tr.get('h'): tr.set('h', str(int(int(tr.get('h')) * s)))
    for gc in el.iter(A + 'gridCol'):
        if gc.get('w'): gc.set('w', str(int(int(gc.get('w')) * s)))

def kind_of(el, name):
    if name == 'Cita': return 'cite'
    if el.tag == MC + 'AlternateContent': return 'eq'
    if el.tag == P + 'pic': return 'pic'
    if el.tag == P + 'graphicFrame': return 'table' if el.find('.//a:tbl', NS) is not None else 'chart'
    if el.tag == P + 'grpSp': return 'group'
    txt = ''.join(t.text or '' for t in el.iter(A + 't'))
    return 'text' if txt.strip() else 'fill'

def set_text(el, text):
    ps = [p for p in el.iter(A + 'p') if not in_fallback(p)]
    tx = ps[0].getparent()
    for p in ps[1:]: p.getparent().remove(p)
    r0 = next((r for r in ps[0].iter(A + 'r') if not in_fallback(r)), None)
    if r0 is None or r0.find(A + 'rPr') is None:               # párrafo vacío: el run toma el formato de fin de párrafo
        end = ps[0].find(A + 'endParaRPr')
        rp = etree.Element(A + 'rPr', end.attrib) if end is not None else etree.Element(A + 'rPr')
        if end is not None:
            for ch in end: rp.append(copy.deepcopy(ch))
        if r0 is None:
            r0 = etree.Element(A + 'r'); etree.SubElement(r0, A + 't')
            (end.addprevious(r0) if end is not None else ps[0].append(r0))
        r0.insert(0, rp)
    plain_par(ps[0], text)

def larga_title(s):
    for el in top_children(s):
        if next(el.iter(P + 'cNvPr')).get('name') == 'Título':
            return ' '.join(ptext(p) for p in el.iter(A + 'p') if not in_fallback(p)).strip()
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip(): return sh.text_frame.text.strip()
    return ''

def notes_of(s): return s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''

# plantillas de la comprimida: encabezado de escena y pie de fuentes
HEAD_T = copy.deepcopy(tagged(old_slides[1])[1])             # encabezado «Toda obra empieza…»
CITE_T = copy.deepcopy(tagged(old_slides[1])[6])             # «Fuentes: …»
CONTENT = (1.62, 4.92)                                       # banda vertical de contenido de una escena (pulgadas)


# ---------------------------------------------------------------- riel lateral de palabras clave
RAIL_X, RAIL_W, RAIL_Y0, RAIL_Y1, RAIL_GAP = 8.3, 1.25, 1.1, 4.95, 0.06
OLD_SLOT_Y0, OLD_SLOT_DY, OLD_SLOT_H = 1.62, 0.84, 0.78
CONTENT_R, HEAD_W = 8.12, 7.72
def in_rail(r): return r is not None and None not in r[:4] and r[0] >= 8.2 and r[2] <= 1.4

def rail_chip_xml(tag, text, color, r, sz):
    x, y, w, h = (int(v * EMU) for v in r)
    w_ = text.rsplit(' ', 1)                                    # sin palabra corta sola al final del renglón
    if len(w_) == 2 and len(w_[1]) <= 3: text = w_[0] + ' ' + w_[1]
    text = text.replace(' → ', ' → ')
    esc = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    return (f'<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f'<p:nvSpPr><p:cNvPr id="{tag}" name="Clave [o{tag}]"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>'
            f'<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val 18000"/></a:avLst></a:prstGeom>'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr>'
            f'<p:txBody><a:bodyPr wrap="square" lIns="45720" tIns="18288" rIns="45720" bIns="18288" anchor="ctr"/><a:lstStyle/>'
            f'<a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="es-ES" sz="{sz}" b="1"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>'
            f'<a:latin typeface="Arial"/></a:rPr><a:t>{esc}</a:t></a:r></a:p></p:txBody></p:sp>')

def apply_rail(scenes, placed, tree, keys):
    """Riel derecho como ayuda memoria: cada palabra clave entra en la escena SIGUIENTE a la que presenta su
    concepto (las de la última escena, en un clic extra de cierre) y queda hasta el final. En cada escena los
    elementos visibles del riel (miniaturas y palabras clave) se apilan desde arriba en orden de entrada;
    si alguno sale, los demás suben. Si el contenido invade el riel, se escala."""
    if not keys: return
    nsc = len(scenes)
    if any(k >= nsc - 1 for k, _ in keys):                     # clic de cierre: misma escena, solo entra el riel
        last = scenes[-1]
        scenes.append({'objs': [dict(o, b=list(o['g']), zf=None, swap=None) for o in last['objs']], 'fx': []})
        for o in scenes[-1]['objs']: o.pop('swap', None)
    enter = [min(k + 1, len(scenes) - 1) for k, _ in keys]
    fixed = lambda o: o['kind'] in ('head', 'cite')
    maxr = max([r[0] + r[2] for sc in scenes for o in sc['objs'] if not fixed(o)
                for r in (o.get('g'), o.get('b')) if r and not in_rail(r)] + [0])
    f = min(1.0, (CONTENT_R - 0.3) / (maxr - 0.3)) if maxr > CONTENT_R else 1.0
    dy = (3.3 * (1 - f)) / 2
    nz = lambda v, fn: None if v is None else fn(v)
    Tc = lambda r: [nz(r[0], lambda v: 0.3 + (v - 0.3) * f), nz(r[1], lambda v: 1.62 + dy + (v - 1.62) * f),
                    nz(r[2], lambda v: v * f), nz(r[3], lambda v: v * f)]
    slot_of = lambda r: max(0, int(round((r[1] - OLD_SLOT_Y0) / OLD_SLOT_DY)))
    first = {}
    for k, sc in enumerate(scenes):
        for o in sc['objs']:
            if in_rail(o['g']): first.setdefault(slot_of(o['g']), k)
    order = sorted([(k, 0, s) for s, k in first.items()] + [(enter[i], 1, i) for i in range(len(keys))])
    rank = {(t, s): j for j, (_, t, s) in enumerate(order)}     # orden de entrada al riel
    # casilleros visibles por escena
    vis = []
    for k, sc in enumerate(scenes):
        v = {('t', slot_of(o['g'])) for o in sc['objs'] if in_rail(o['g'])} | {('k', i) for i in range(len(keys)) if enter[i] <= k}
        vis.append(sorted(v, key=lambda e: rank[(0 if e[0] == 't' else 1, e[1])]))
    N = max(len(v) for v in vis) or 1
    h = min(OLD_SLOT_H, (RAIL_Y1 - RAIL_Y0 - RAIL_GAP * (N - 1)) / N)
    top = lambda j: RAIL_Y0 + j * (h + RAIL_GAP)
    fs = h / OLD_SLOT_H; cx = RAIL_X + RAIL_W / 2
    def Tr(r, k):
        s = slot_of(r); y0 = OLD_SLOT_Y0 + s * OLD_SLOT_DY; j = vis[k].index(('t', s))
        return [cx + (r[0] - cx) * fs, top(j) + (r[1] - y0) * fs, r[2] * fs, r[3] * fs]
    def Tr_any(r):                                              # rects base: casillero de primera entrada
        s = slot_of(r); k = first.get(s, 0); return Tr(r, k) if ('t', s) in vis[k] else r
    def T(r, o, k):
        if r is None: return None
        if in_rail(r): return Tr(r, k) if k is not None else Tr_any(r)
        if o['kind'] == 'head': return [r[0], r[1], min(r[2], HEAD_W), r[3]]
        return r if fixed(o) else Tc(r)
    done = set()
    for k, sc in enumerate(scenes):
        for o in sc['objs']:
            o['g'] = T(o['g'], o, k)
            for key in ('b', 'zf'):
                if o.get(key): o[key] = T(o[key], o, None)
            t = o['tag']
            if t not in done and t in placed:
                done.add(t); el = placed[t]
                r0 = rect_in(el); r1 = T(r0, o, None); set_rect(el, r1)
                if not in_rail(r0) and not fixed(o) and f < 0.999: scale_fonts(el, f)
        for fx in sc.get('fx', []):
            if fx.get('to'): fx['to'] = Tr_any(fx['to']) if in_rail(fx['to']) else Tc(fx['to'])
    sz = 1000 if h >= 0.42 else 900
    for i, (k, text) in enumerate(keys):
        t = new_tag(); kk = enter[i]
        r = [RAIL_X, top(vis[kk].index(('k', i))), RAIL_W, h]
        chip = etree.fromstring(rail_chip_xml(t, text, SPEC.PALETTE[i % len(SPEC.PALETTE)], r, sz))
        tree.append(chip); placed[t] = chip
        for k2 in range(kk, len(scenes)):
            g = [RAIL_X, top(vis[k2].index(('k', i))), RAIL_W, h]
            scenes[k2]['objs'].append({'tag': t, 'kind': 'text', 'g': g, 'b': list(r), 'zf': None})

# ---------------------------------------------------------------- construcción
next_tag = [5000]
def new_tag(): next_tag[0] += 1; return next_tag[0]
log = []
built = []                                                    # (slide, scenes, notes)

for S in SPEC.SLIDES:
    host = old_slides[S['host'] - 1]
    slide = comp.slides.add_slide(host.slide_layout)
    tree = slide._element.find('.//p:spTree', NS)
    for ph in list(slide.placeholders): ph._element.getparent().remove(ph._element)
    hbg = host._element.find('p:cSld/p:bg', NS)                # fondo propio (portada, preguntas)
    if hbg is not None: slide._element.find('p:cSld', NS).insert(0, copy_el(hbg, host.part, slide.part, False))
    for el in top_children(host):                             # fijos: título general, número, portada…
        if (tag_of(el) is None or not S['scenes']) and not S.get('no_static'):
            tree.append(copy_el(el, host.part, slide.part, False))
    if S.get('title') is not None:
        for el in top_children(slide):
            if next(el.iter(P + 'cNvPr')).get('name') == 'Título' and tag_of(el) is None: set_text(el, S['title'])
    placed = {}                                               # tag -> elemento en la slide nueva
    fresh = []                                                # elementos tomados de la larga v37: el textmap no los toca
    scenes, notes = [], []

    def place(tag, el_src, src_part, cross=False):
        if tag in placed: return
        el = copy_el(el_src, src_part, slide.part, cross)
        num = next((e for e in top_children(slide) if e.find('.//p:nvPr/p:ph', NS) is not None), None)
        (num.addprevious(el) if num is not None else tree.append(el)); placed[tag] = el

    for sc in S['scenes']:
        if sc[0] == 'old':
            _, ci, si = sc[:3]; opts = sc[3] if len(sc) > 3 else {}
            osl = old_slides[ci - 1]; tg = tagged(osl)
            o_sc = copy.deepcopy(CH[str(ci)]['scenes'][si])
            o_sc['objs'] = [o for o in o_sc['objs'] if o['tag'] not in opts.get('drop', [])]
            prev_orig = {o['tag'] for o in CH[str(ci)]['scenes'][si - 1]['objs']} if si else set()
            o_sc['objs'] = [o for o in o_sc['objs'] if not (in_rail(o['g']) and o['tag'] not in placed)]
            for o in o_sc['objs']:
                is_new = o["tag"] not in placed
                place(o['tag'], tg[o['tag']], osl.part)
                if is_new and o["tag"] in prev_orig:         # venía de una escena que aquí no está: entra ya en su lugar
                    set_rect(placed[o['tag']], o['g']); o['b'] = list(o['g']); o.pop('swap', None); o['zf'] = None
            for t, img in opts.get('img', {}).items():         # cambiar la imagen de un objeto
                pic = placed[int(t)]; blip = pic.find('.//a:blip', NS)
                _, rid = slide.part.get_or_add_image_part(img); blip.set(R + 'embed', rid)
            for t, txt in opts.get('head', {}).items(): set_text(placed[int(t)], txt)
            ln = opts.get('larga')
            if ln:                                              # pie de fuentes de la v37
                lc = next((e for e in top_children(L[ln - 1]) if next(e.iter(P + 'cNvPr')).get('name') == 'Cita'), None)
                lt = ptext(lc.find('.//a:p', NS)).strip() if lc is not None else ''
                for o in o_sc['objs']:
                    if o['kind'] == 'cite' and o['tag'] in placed: set_text(placed[o['tag']], ("Fuentes: " + lt) if lt else '')
            scenes.append(o_sc)
            ln = opts.get('larga')
            if ln: notes.append(f"[Slide larga {ln}: {larga_title(L[ln - 1])}] {notes_of(L[ln - 1])}".strip())
        elif sc[0] == 'new':
            ln = sc[1]; opts = sc[2] if len(sc) > 2 else {}
            ls = L[ln - 1]; objs = []
            # encabezado
            ht = new_tag(); h = copy.deepcopy(HEAD_T); set_tag(h, ht); set_text(h, opts.get('head', larga_title(ls)))
            tree.append(h); placed[ht] = h; objs.append({'tag': ht, 'kind': 'head', 'g': rect_in(h), 'b': rect_in(h), 'zf': None})
            # contenido: banda [y0, y1] de la larga -> banda de contenido de la comprimida
            els = [e for e in top_children(ls) if e.find('.//p:nvPr/p:ph', NS) is None
                   and next(e.iter(P + 'cNvPr')).get('name') not in ('Título', 'Cita')]
            if els:
                rs = [rect_in(e) for e in els]
                y0 = min(r[1] for r in rs); y1 = max(r[1] + r[3] for r in rs)
                x0 = min(r[0] for r in rs); x1 = max(r[0] + r[2] for r in rs)
                s = min(1.0, (CONTENT[1] - CONTENT[0]) / max(y1 - y0, 0.01), 9.2 / max(x1 - x0, 0.01))
                cx = (x0 + x1) / 2
                for e, r in zip(els, rs):
                    t = new_tag(); el = copy_el(e, ls.part, slide.part, True); set_tag(el, t)
                    nr = [5.0 + (r[0] - cx) * s, CONTENT[0] + (r[1] - y0) * s, r[2] * s, r[3] * s]
                    set_rect(el, nr)
                    if s < 0.999: scale_fonts(el, s)
                    tree.append(el); placed[t] = el; fresh.append(el)
                    objs.append({'tag': t, 'kind': kind_of(e, next(e.iter(P + 'cNvPr')).get('name')), 'g': nr, 'b': nr, 'zf': None})
            # fuentes
            cita = next((e for e in top_children(ls) if next(e.iter(P + 'cNvPr')).get('name') == 'Cita'), None)
            ctext = ptext(cita.find('.//a:p', NS)).strip() if cita is not None else ''
            if ctext:
                ct = new_tag(); c = copy.deepcopy(CITE_T); set_tag(c, ct); set_text(c, "Fuentes: " + ctext)
                tree.append(c); placed[ct] = c; objs.append({'tag': ct, 'kind': 'cite', 'g': rect_in(c), 'b': rect_in(c), 'zf': None})
            if scenes:                                          # el riel de miniaturas sigue visible
                for o in scenes[-1]['objs']:
                    if in_rail(o['g']): objs.append({'tag': o['tag'], 'kind': o['kind'], 'g': list(o['g']), 'b': list(o['g']), 'zf': None})
            scenes.append({'objs': objs, 'fx': []})
            notes.append(f"[Slide larga {ln}: {larga_title(ls)}] {notes_of(ls)}".strip())
        elif sc[0] == 'keys':
            chips = sc[1]; objs = []
            ht = new_tag(); h = copy.deepcopy(HEAD_T); set_tag(h, ht); set_text(h, sc[2] if len(sc) > 2 else "En resumen")
            tree.append(h); placed[ht] = h; objs.append({'tag': ht, 'kind': 'head', 'g': rect_in(h), 'b': rect_in(h), 'zf': None})
            n = len(chips); cols = 2 if n <= 4 else 3; rows = (n + cols - 1) // cols
            W, Hh, gx, gy = 9.0, 3.1, 0.25, 0.22
            cw = (W - gx * (cols - 1)) / cols; chh = min(1.2, (Hh - gy * (rows - 1)) / rows)
            top = 1.75 + (Hh - (rows * chh + (rows - 1) * gy)) / 2
            pal = SPEC.PALETTE
            for i, (kw, det) in enumerate(chips):
                r_, c_ = divmod(i, cols)
                row_n = min(cols, n - r_ * cols); left0 = 0.5 + (W - (row_n * cw + (row_n - 1) * gx)) / 2
                rect = [left0 + c_ * (cw + gx), top + r_ * (chh + gy), cw, chh]
                t = new_tag(); col = pal[i % len(pal)]
                chip = etree.fromstring(SPEC.chip_xml(t, kw, det, col, rect, EMU))
                tree.append(chip); placed[t] = chip
                objs.append({'tag': t, 'kind': 'text', 'g': rect, 'b': rect, 'zf': None})
            scenes.append({'objs': objs, 'fx': []})
    apply_rail(scenes, placed, tree, SPEC.KEYS.get(S.get('title') or '', []))
    # continuidad: un objeto que deja la escena y vuelve más tarde se clona (anim3 no reutiliza instancias)
    present_prev, seen, alias = set(), set(), {}
    for k, sc in enumerate(scenes):
        cur = set()
        for o in sc['objs']:
            t = o['tag']
            if t in alias and alias[t][1] == k - 1: o['tag'] = alias[t][0]; alias[t] = (o['tag'], k)
            elif t in seen and t not in present_prev:
                nt = new_tag(); el = copy.deepcopy(placed[t]); set_tag(el, nt)
                g = o['g']; set_rect(el, g); tree.append(el); placed[nt] = el
                o['tag'] = nt; o['b'] = list(g); o.pop('swap', None); o['zf'] = None; alias[t] = (nt, k)
            seen.add(o['tag']); cur.add(o['tag'])
        present_prev = cur
    # un solo movimiento por objeto: si hay un segundo, se clona en el nuevo lugar
    state = {}
    for k, sc in enumerate(scenes):
        for o in sc['objs']:
            t = o['tag']
            if t in state and anim3.rect_diff(state[t]['r'], o['g']) > anim3.EPS:
                if state[t]['mv']:
                    nt = new_tag(); el = copy.deepcopy(placed[t]); set_tag(el, nt); set_rect(el, o['g'])
                    tree.append(el); placed[nt] = el
                    for sc2 in scenes[k:]:
                        for o2 in sc2['objs']:
                            if o2['tag'] == t: o2['tag'] = nt; o2['b'] = list(o['g'])
                    state[nt] = {'r': list(o['g']), 'mv': False}; continue
                state[t]['mv'] = True; state[t]['r'] = list(o['g'])
            elif t not in state:
                state[t] = {'r': list(o.get('b') or o['g']), 'mv': False}
    # fx que apuntan a objetos que ya no salen en la escena siguiente: se descartan
    for k, sc in enumerate(scenes):
        prev = {o['tag'] for o in scenes[k - 1]['objs']} if k else set()
        sc['fx'] = [f for f in sc.get('fx', []) if f['tag'] in prev]
    movers, first_r = [], {}                                  # lo que se mueve o va al riel queda delante
    for sc in scenes:
        for o in sc['objs']:
            t = o['tag']
            r0 = first_r.setdefault(t, o.get('b') or o['g'])
            if t in placed and t not in movers and anim3.rect_diff(r0, o['g']) > anim3.EPS: movers.append(t)
    for t in movers:
        el = placed[t]
        if el.getparent() is tree: tree.remove(el); tree.append(el)
    S['_fresh'] = fresh
    built.append((slide, scenes, notes, S))
    log.append(f"{len(built):2d} {S.get('title') or ''} — {len(scenes)} escenas")

# ---------------------------------------------------------------- textos v29 -> v37
pairs = {k: v for k, v in TMAP['pairs']}; XMLS = TMAP.get('xml', {})
deleted = {d for d in TMAP['deleted'] if len(d) > 25}
nrep = ndel = 0
for slide, scenes, notes, S in built:
    fids = {id(e) for e in S['_fresh']}                       # los elementos siguen vivos en S: sus id son estables
    for p in list(slide._element.iter(A + 'p')):
        if in_fallback(p) or any(id(a) in fids for a in p.iterancestors()): continue
        t = ptext(p).strip()
        if t.isdigit(): continue                              # número de slide
        if t in XMLS and ('a14:m' in XMLS[t] or any(True for _ in p.iter(M + 't'))):   # con fórmulas: se injerta el párrafo de la v37
            if True:
                np_ = etree.fromstring(XMLS[t])
                for fb in np_.iter(MC + 'Fallback'):          # el respaldo de la larga apunta a partes de otro paquete
                    ch_ = fb.getprevious(); txt = ''.join(t_.text or '' for t_ in ch_.iter(M + 't'))
                    for ch in list(fb): fb.remove(ch)
                    r_ = etree.SubElement(fb, A + 'r'); etree.SubElement(r_, A + 'rPr', lang='es-ES'); etree.SubElement(r_, A + 't').text = txt or ' '
                szs = [r.get('sz') for r in p.iter(A + 'rPr') if r.get('sz')]
                for ch in list(p):
                    if ch.tag not in (A + 'pPr', A + 'endParaRPr'): p.remove(ch)
                end = p.find(A + 'endParaRPr')
                for ch in list(np_):
                    if ch.tag in (A + 'pPr', A + 'endParaRPr'): continue
                    if szs:
                        for r in ch.iter(A + 'rPr'): r.set('sz', szs[0])
                    (end.addprevious(ch) if end is not None else p.append(ch))
                nrep += 1
        elif t in pairs:
            plain_par(p, [(x, k) for x, k in pairs[t]]); nrep += 1
        elif t in deleted and len([q for q in p.getparent().findall(A + 'p')]) > 1:
            p.getparent().remove(p); ndel += 1
    for p in list(slide._element.iter(A + 'p')):              # pie de fuentes vacío
        if not in_fallback(p) and ptext(p).strip() in ('Fuentes:', 'Fuente:'): plain_par(p, '')
    for mt in slide._element.iter(M + 't'):
        if mt.text and '𝜇' in mt.text: mt.text = mt.text.replace('𝜇', '𝐺')

# ---------------------------------------------------------------- notas, ids, coreografía
for i, (slide, scenes, notes, S) in enumerate(built, 1):
    if notes: slide.notes_slide.notes_text_frame.text = "\n\n".join(notes)
    xml = etree.tostring(slide._element, encoding='unicode')
    new = anim3.normalize_ids(xml)
    if new != xml:
        root = etree.fromstring(new.encode('utf-8'))
        old_tree = slide._element.find('p:cSld', NS); old_tree.getparent().replace(old_tree, root.find('p:cSld', NS))
    for t in slide._element.findall('p:timing', NS): slide._element.remove(t)
    for t in slide._element.findall('p:transition', NS): slide._element.remove(t)
    if len(scenes) > 1:
        xml = etree.tostring(slide._element, encoding='unicode')
        spids = {int(m.group(2)): int(m.group(1)) for m in re.finditer(r'<p:cNvPr id="(\d+)" name="[^"]*\[o(\d+)\]"', xml)}
        paras = {}
        for mm in re.finditer(r'<p:sp>.*?</p:sp>', xml, re.S):
            seg = mm.group(0); tg = re.search(r'name="[^"]*\[o(\d+)\]"', seg)
            if tg:
                tb = re.search(r'<p:txBody>.*?</p:txBody>', seg, re.S)
                paras[int(tg.group(1))] = len(re.findall(r'<a:p>|<a:p ', tb.group(0))) if tb else 1
        kinds = {o['tag']: o['kind'] for sc in scenes for o in sc['objs']}
        for sc in scenes:
            for o in sc['objs']:
                for key in ('g', 'b', 'zf'):
                    if o.get(key): o[key] = [0.0 if v is None else v for v in o[key]]
        miss = {o['tag'] for sc in scenes for o in sc['objs']} - set(spids)
        assert not miss, (i, 'tags sin shape', sorted(miss)[:5])
        tim, nc = anim3.build_timing(scenes, spids, kinds, paras)
        tim = tim.replace('<p:timing>', '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">', 1)
        tnode = etree.fromstring('<x xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
                                 '<p:transition spd="med"><p:fade/></p:transition></x>')[0]
        cmo = slide._element.find('p:clrMapOvr', NS)
        anchor = cmo if cmo is not None else slide._element.find('p:cSld', NS)
        anchor.addnext(tnode); tnode.addnext(etree.fromstring(tim))

# ---------------------------------------------------------------- quitar las slides viejas
sld = comp.slides._sldIdLst
olds = {s.part for s in old_slides}
for sid in list(sld):
    rid = sid.get(R + 'id')
    if comp.part.related_part(rid) in olds: comp.part.drop_rel(rid); sld.remove(sid)
comp.save(out_path)
json.dump({str(i): {'scenes': sc} for i, (_, sc, _, _) in enumerate(built, 1)}, open(out_path + '.choreo.json', 'w', encoding='utf-8'), ensure_ascii=False)
print('\n'.join(log)); print(f"textos reemplazados: {nrep}, párrafos borrados: {ndel}; slides: {len(built)}")
