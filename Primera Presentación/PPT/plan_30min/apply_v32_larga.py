"""v32 larga (después de restructure_v32_larga.ps1)."""
import sys, copy, os
import pptx
from pptx.util import Emu
from pptx_helpers_v30 import NS, A, P, ptext, top_shapes, unwrap, plain_par, body_par, set_title, set_notes

deck = sys.argv[1]
HERE = os.path.dirname(os.path.abspath(__file__)); FIG = os.path.join(HERE, "figuras_es_v25")
EXC = os.path.join(HERE, "..", "..", "latex-15p-review", "figuras", "excitacion")
FTM_IMG = r"C:\Github\FTM-test\experiments\scope\FTM mediciones\retrasos.png"
prs = pptx.Presentation(deck)
S = lambda n: prs.slides[n - 1]
title = lambda n: ptext(top_shapes(S(n), "Título")[1][0].find('.//a:p', NS))
shape = lambda n, nm: next(sh for sh in S(n).shapes if sh.name == nm)

def check(n, start):
    assert title(n).startswith(start), f"slide {n}: «{title(n)[:60]}» no empieza con «{start}»"

def set_text(n, name, text):
    tf = shape(n, name).text_frame; p = tf.paragraphs[0]
    if p.runs: p.runs[0].text = text; [r._r.getparent().remove(r._r) for r in p.runs[1:]]
    else: p.add_run().text = text

def fit_picture(n, img, box, name):
    x, y, w, h = box
    pic = S(n).shapes.add_picture(img, x, y)
    r = min(w / pic.width, h / pic.height)
    pic.width, pic.height = int(pic.width * r), int(pic.height * r)
    pic.left, pic.top = int(x + (w - pic.width) / 2), int(y + (h - pic.height) / 2); pic.name = name
    return pic

def bullets(n, lines):
    sp = unwrap(S(n), "Cuerpo"); tx = sp.find('.//p:txBody', NS); ps = tx.findall('a:p', NS)
    for p in ps[1:]: tx.remove(p)
    plain_par(ps[0], lines[0])
    for line in lines[1:]:
        q = copy.deepcopy(ps[0]); plain_par(q, line); tx.append(q)

# 62: la compensación inversa satura (coherente con la 63 y con la 66)
check(62, "Existen cuatro familias")
set_text(62, "Text 8", "Se cancela la dinámica del sensor para reconstruir la baja frecuencia. Amplifica el ruido "
                       "de baja frecuencia y el offset, y lleva fácilmente a la saturación.")

# 64: carga activa Z_L
check(64, "Resistencia negativa y balance")
plain_par(body_par(S(64), "Cuerpo", "Cambiar la resistencia de carga")[1],
          [("Una carga activa ", "n"), ("Z", "i"), ("L", "s"), (", como una resistencia negativa, modifica el amortiguamiento electromecánico", "n")])

# 79: HAMMER sin pie de fuente (el 086D20 es ajeno)
check(79, "El nodo HAMMER")
cita = copy.deepcopy(shape(80, "Cita")._element)
ids = [int(e.get('id')) for e in S(79)._element.iter(P + 'cNvPr')]
next(cita.iter(P + 'cNvPr')).set('id', str(max(ids) + 1))
S(79).shapes._spTree.append(cita)
set_text(79, "Cita", "PCB Piezotronics (2025).")

# 80: figura delta-sigma con la flecha sobre la curva
check(80, "Selección del tipo de ADC")
old = shape(80, "Conformación del ruido")
pic = fit_picture(80, os.path.join(FIG, "fig_delta_sigma_es.png"), (old.left, old.top, old.width, old.height), "Conformación del ruido")
old._element.addprevious(pic._element); old._element.getparent().remove(old._element)

# 95: tabla de alternativas con fuentes; UWB pasa a FTM
check(95, "Se eligió ESP-NOW")
tbl = next(sh for sh in S(95).shapes if sh.has_table).table
cell = lambda r, c, v: plain_par(tbl.cell(r, c)._tc.find('.//a:p', NS), v)
cell(4, 2, "Referencia absoluta por PPS")
cell(4, 3, "La usan equipos comerciales, como el Atom-3C. Contras: costo y cobertura")
cell(6, 0, "Wi-Fi FTM (ESP32-S3)"); cell(6, 2, "≲ 150 ns entre esclavos (preliminar)"); cell(6, 3, "Trabajo en curso")
set_text(95, "Cita", "Wi-Fi: IEEE Std 802.11-2020. LoRa: Semtech (2013), AN1200.13. GPS: Geometrics (2026), Atom-3C. "
                     "FTM: abbbe (s. f.), fts.")

# 114: trabajo en curso, FTM (copia de la 61)
check(114, "Por qué importa la pérdida")
set_title(S(114), "Trabajo en curso: sincronización por Wi-Fi FTM")
old = shape(114, "Geófonos 10 y 4,5 Hz")
pic = fit_picture(114, FTM_IMG, (old.left, old.top, old.width, old.height), "Osciloscopio FTM")
old._element.addprevious(pic._element); old._element.getparent().remove(old._element)
bullets(114, ["FTM (Fine Timing Measurement, IEEE 802.11) mide tiempos de vuelo entre ESP32-S3",
              "Con FTS, cada esclavo ajusta su reloj al del maestro: 85 de 85 sesiones completas",
              "Entre esclavos, los flancos caen dentro de unos 150 ns (preliminar, canales sin identificar)",
              "También estima distancias entre nodos, todavía sin precisión estable"])
set_text(114, "Text 10", "")
set_text(114, "Cita", "FTS: abbbe (s. f.), fts. Osciloscopio: GPIO4 de maestro y esclavos, 50 ns/div.")
set_notes(S(114), "Para la sincronización, la línea que estoy probando no es UWB sino Wi-Fi FTM, que ya está en los ESP32-S3. "
                  "FTM mide tiempos de vuelo entre nodos. Sobre el proyecto FTS, cada esclavo ajusta su reloj al del maestro. "
                  "En la prueba de concepto, las 85 sesiones terminaron completas. En el osciloscopio, los flancos de los esclavos "
                  "caen dentro de unos 150 nanosegundos entre sí, tres órdenes por debajo de los 200 a 400 microsegundos de ESP-NOW. "
                  "Es preliminar: los canales no están identificados y falta medir con el pulso de hardware. "
                  "FTM además estima distancias, que servirían para controlar la geometría medida con cinta.")

# 115: trabajo en curso, martillete de leva (copia de la 61)
check(115, "Por qué importa la pérdida")
set_title(S(115), "Trabajo en curso: martillete de leva inspirado en Da Vinci")
old = shape(115, "Geófonos 10 y 4,5 Hz")
x, y, w, h = old.left, old.top, old.width, old.height
p1 = fit_picture(115, os.path.join(EXC, "concepto_martillete_da_vinci.png"), (x, y, int(w * 0.40), h), "Concepto martillete")
p2 = fit_picture(115, os.path.join(EXC, "estado_actual_martillete.jpg"), (x + int(w * 0.43), y, int(w * 0.57), h), "Estado martillete")
for q in (p1, p2): old._element.addprevious(q._element)
old._element.getparent().remove(old._element)
bullets(115, ["Una leva eleva y libera la masa: trenes de impulsos con repetición controlada",
              "Ya están armados el bastidor, el motor y el martillo",
              "Falta integrar transmisión, guiado, control y protecciones, y ensayar",
              "Objetivo: una fuente más repetible que el golpe manual"])
set_text(115, "Text 10", "")
set_text(115, "Cita", "")
set_notes(S(115), "La otra línea en curso es la fuente. El golpe manual varía de un impacto a otro, y eso obliga a promediar "
                  "muchos golpes. Armamos un martillete con un mecanismo de leva, inspirado en el de Da Vinci, que eleva y libera "
                  "la masa de forma repetida. Ya están el bastidor, el motor y el martillo. Falta integrar la transmisión, el "
                  "guiado, el control y las protecciones, y recién después ensayarlo.")

# 123: divisor de anexo (copia de «Preguntas»)
set_text(123, "Text 0", "Anexo"); set_text(123, "Text 1", "Material de respaldo")
set_notes(S(123), "")

# referencias: entran las fuentes de la tabla y de FTM; salen las de UWB si ya no se citan
NEW = ["abbbe (s. f.). fts [repositorio de software]. github.com/abbbe/fts.",
       "Geometrics (2026). Atom-3C Seismograph. geometrics.com/product/atom-3c.",
       "IEEE (2020). IEEE Std 802.11-2020: Wireless LAN Medium Access Control (MAC) and Physical Layer (PHY) Specifications.",
       "Semtech Corporation (2013). SX1272/3/6/7/8: LoRa Modem Designer's Guide, AN1200.13, rev. 1."]
body = " ".join(ptext(p) for sl in prs.slides for p in sl._element.iter(A + 'p')
                if not any(sh.name == "Referencias" for sh in sl.shapes) or True)
boxes = [sh for sl in prs.slides for sh in sl.shapes if sh.name == "Referencias"]
cited_text = " ".join(ptext(p) for sl in prs.slides if not any(sh.name == "Referencias" for sh in sl.shapes)
                      for p in sl._element.iter(A + 'p'))
removed = []
for b in boxes:
    for p in list(b.text_frame.paragraphs):
        t = p.text
        for key in ("Korsimaa", "Xue, B.", "Pérez-Solano"):
            if t.startswith(key) and key.split(",")[0] not in cited_text:
                p._p.getparent().remove(p._p); removed.append(key)
tpl = boxes[0].text_frame.paragraphs[0]._p
for text in NEW:
    e = copy.deepcopy(tpl); rs = e.findall(A + 'r'); rs[0].find(A + 't').text = text
    for r in rs[1:]: e.remove(r)
    boxes[-1].text_frame._txBody.append(e)

prs.save(deck)
print("ok; referencias quitadas:", removed)
