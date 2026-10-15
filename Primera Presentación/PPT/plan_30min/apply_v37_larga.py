"""v37 larga: balance de fuerza en dos slides con figuras de X. Zhang et al. (2020).
64: modelo (fig. 2). 65 (copia de la 64): respuesta medida (fig. 9) con los resultados del artículo.
Datos verificados en el PDF: SM-6 de 4,5 Hz; −3 dB de 4,5 Hz a 0,16 Hz; límite superior 120 Hz por modos espurios;
4,3 ng/√Hz a 1 Hz (resumen y sec. 5)."""
import sys, os, copy
import pptx
from pptx.util import Pt
from pptx_helpers_v30 import NS, A, ptext, top_shapes, unwrap, plain_par, set_title, set_notes

deck = sys.argv[1]
FIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figuras_es_v25")
prs = pptx.Presentation(deck)
S = lambda n: prs.slides[n - 1]
for n in (64, 65): assert ptext(top_shapes(S(n), "Título")[1][0].find('.//a:p', NS)) == "Balance de fuerza"

def set_text(s, name, text):
    sh = next(x for x in s.shapes if x.name == name); p = sh.text_frame.paragraphs[0]
    if p.runs: p.runs[0].text = text; [r._r.getparent().remove(r._r) for r in p.runs[1:]]
    else: p.add_run().text = text
    return sh

# ---- 64: el modelo; el mensaje pasa a la 65 ----
s = S(64)
cue = unwrap(s, "Cuerpo"); plain_par(cue.find('.//a:p', NS), "La misma bobina mide el movimiento y, con un lazo PID, aplica la fuerza de balance")
msg = set_text(s, "Text 4", "")
set_text(s, "Cita", "Imagen: X. Zhang et al. (2020), Sens. Actuators A 312, 112074, fig. 2, traducida. Ver también Z. Zhang et al. (2025).")
set_notes(s, "La segunda técnica activa es el balance de fuerza. La figura es el modelo de Zhang y colaboradores, de 2020: "
             "la misma bobina del geófono mide y actúa. La tensión de la bobina pasa por un modulador sigma-delta de cuarto "
             "orden; el tren de bits vuelve por un DAC de un bit y un controlador PID, y la bobina aplica una fuerza de Lorentz "
             "que se opone al movimiento. En la slide siguiente, qué consiguen.")

# ---- 65: la respuesta medida ----
s = S(65)
set_title(s, "Balance de fuerza: respuesta medida")
sh = {x.name: x for x in s.shapes}
old = sh["Lazo de balance de fuerza"]; tit = sh["Título"]
top = tit.top + tit.height + Pt(4); bottom = Pt(318)
pic = s.shapes.add_picture(os.path.join(FIG, "zhang2020_fig9_traducida.png"), Pt(36), top)
r = (bottom - top) / pic.height; pic.width, pic.height = int(pic.width * r), int(pic.height * r)
pic.name = "Respuesta medida Zhang 2020"
old._element.addprevious(pic._element); old._element.getparent().remove(old._element)
cue = unwrap(s, "Cuerpo"); tx = cue.find('.//p:txBody', NS); p0 = tx.find('a:p', NS)
from lxml import etree
pPr = p0.find('a:pPr', NS)                      # vuelve a ser lista con viñetas
pPr.set('marL', '342900'); pPr.set('indent', '-342900')
bn = pPr.find('a:buNone', NS)
if bn is not None:
    ch = etree.Element(A + 'buChar'); ch.set('char', '•'); bn.addprevious(ch); pPr.remove(bn)
lines = ["Con un geófono SM-6 de 4,5 Hz, el punto de −3 dB baja de 4,5 Hz a 0,16 Hz",
         "Arriba se limita a 120 Hz por los modos espurios del geófono",
         "Ruido equivalente: 4,3 ng/√Hz a 1 Hz"]
plain_par(p0, lines[0])
for line in lines[1:]:
    q = copy.deepcopy(p0); plain_par(q, line); tx.append(q)
for sp_ in cue.iter(A + 'spcPts'): sp_.set('val', '600')
x0 = pic.left + pic.width + Pt(18)
off = cue.find('.//a:xfrm/a:off', NS); ext = cue.find('.//a:xfrm/a:ext', NS)
off.set('x', str(x0)); off.set('y', str(top + Pt(10))); ext.set('cx', str(Pt(684) - x0)); ext.set('cy', str(bottom - top - Pt(10)))
msg = set_text(s, "Text 4", "Funciona, pero su complejidad pesa en un nodo multicanal barato y replicable.")
msg.top, msg.left, msg.width, msg.height = Pt(326), Pt(36), Pt(648), Pt(26)
set_text(s, "Cita", "Imagen: X. Zhang et al. (2020), Sens. Actuators A 312, 112074, fig. 9, traducida.")
set_notes(s, "Esto es lo que consiguen. En azul, su geófono digital; en negro, un SM-6 convencional de 4,5 hertz. "
             "Con el lazo, el punto de menos 3 decibeles baja de 4,5 hertz a 0,16 hertz. Arriba lo limitan a 120 hertz para "
             "no excitar los modos espurios del geófono, que se ven como la caída cerca de 150 hertz. El ruido equivalente es "
             "de 4,3 nanogravedades por raíz de hertz a 1 hertz. Funciona muy bien, pero es un lazo analógico y digital pegado "
             "al transductor, con FPGA, y para un nodo multicanal barato y replicable ese costo de complejidad pesa.")

# ---- referencia con los autores completos y el DOI ----
for sl in prs.slides:
    for x in sl.shapes:
        if x.name != "Referencias": continue
        for p in x.text_frame.paragraphs:
            if p.text.startswith("Zhang, X., Wei, X."):
                p.runs[0].text = ("Zhang, X., Wei, X., Wang, T., Li, X., Xiao, W. y Wu, Z. (2020). A digital low-frequency "
                                  "geophone based on 4th-order sigma-delta modulator and single-coil velocity feedback. "
                                  "Sensors and Actuators A: Physical, 312, 112074. doi:10.1016/j.sna.2020.112074.")
                for r_ in p.runs[1:]: r_._r.getparent().remove(r_._r)

prs.save(deck); print("ok")
