"""v34 larga (sobre la v33 reordenada por Elías): diagrama de bloques de control en la 68 y texto de la 48."""
import sys, os
import pptx
from pptx.util import Pt
from pptx_helpers_v30 import NS, ptext, top_shapes

deck = sys.argv[1]
FIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figuras_es_v25")
prs = pptx.Presentation(deck)
S = lambda n: prs.slides[n - 1]
title = lambda n: ptext(top_shapes(S(n), "Título")[1][0].find('.//a:p', NS))

# 68: el diagrama de cajas de colores se reemplaza por uno de control clásico
assert title(68).startswith("Implementación del AFE"), title(68)
s = S(68)
keep = {"Título", "Cuerpo", "Cita", "Slide Number Placeholder 0"}
tree = s._element.find('.//p:spTree', NS)
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
for ch in list(tree):
    names = [e.get('name') for e in ch.iter(P + 'cNvPr')]
    if names and names[0] and not (set(names) & keep): tree.remove(ch)
tit = next(sh for sh in s.shapes if sh.name == "Título")
cue_sp = top_shapes(s, "Cuerpo")[1][0]
class _Box: pass
cue = _Box(); cue.top = int(cue_sp.find('.//a:xfrm/a:off', NS).get('y'))
top = tit.top + tit.height + Pt(6); h = cue.top - Pt(8) - top; w = Pt(648)
pic = s.shapes.add_picture(os.path.join(FIG, "fig_afe_bloques_control_es.png"), Pt(36), top)
r = min(w / pic.width, h / pic.height)
pic.width, pic.height = int(pic.width * r), int(pic.height * r)
pic.left = Pt(36) + (w - pic.width) // 2; pic.top = top + (h - pic.height) // 2
pic.name = "Diagrama de bloques del AFE"

# 48: «Otras mediciones» acotan los parámetros supuestos, no se extrapolan
assert title(48).startswith("Límites de la inversión"), title(48)
p = next(sh for sh in S(48).shapes if sh.name == "Text 16").text_frame.paragraphs[0]
assert "Extrapolarla" in p.text, p.text
p.runs[0].text = "Acotan los parámetros supuestos, como ρ, a rangos más realistas y reducen la incertidumbre"
for r_ in p.runs[1:]: r_._r.getparent().remove(r_._r)

prs.save(deck); print("ok")
