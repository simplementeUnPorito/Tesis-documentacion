"""v31 larga (después de restructure_v31_larga.ps1): notación G, slides nuevas 35 (reflexión) y 39 (Rayleigh),
títulos de «primer modelo» y «modelo estratificado», tabla de sensores, figuras nuevas y textos puntuales."""
import sys, copy, os
import pptx
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx_helpers_v30 import NS, A, M, P, ptext, top_shapes, unwrap, plain_par, body_par, set_title, set_notes

deck = sys.argv[1]; FIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figuras_es_v25")
prs = pptx.Presentation(deck)
S = lambda n: prs.slides[n - 1]
title = lambda n: ptext(top_shapes(S(n), "Título")[1][0].find('.//a:p', NS))

def check(n, start):
    assert title(n).startswith(start), f"slide {n}: «{title(n)[:60]}» no empieza con «{start}»"

def shape(n, name):
    return next(sh for sh in S(n).shapes if sh.name == name)

def swap_picture(n, name, img, new_name, box=None):
    """Reemplaza una imagen por otra, ajustada (sin deformar) y centrada en la caja original o en box=(x,y,w,h)."""
    old = shape(n, name)
    x, y, w, h = box or (old.left, old.top, old.width, old.height)
    pic = S(n).shapes.add_picture(img, x, y)
    r = min(w / pic.width, h / pic.height)
    pic.width, pic.height = int(pic.width * r), int(pic.height * r)
    pic.left, pic.top = int(x + (w - pic.width) / 2), int(y + (h - pic.height) / 2)
    pic.name = new_name
    old._element.addprevious(pic._element); old._element.getparent().remove(old._element)
    return pic

def remove(n, name):
    cont, _ = top_shapes(S(n), name); cont.getparent().remove(cont)

def set_cita(n, text):
    tf = shape(n, "Cita").text_frame; p = tf.paragraphs[0]
    if p.runs: p.runs[0].text = text; [r._r.getparent().remove(r._r) for r in p.runs[1:]]
    else: p.add_run().text = text

def math_replace(n, old, new):
    k = 0
    for t in S(n)._element.iter(M + 't'):
        if t.text and old in t.text: t.text = t.text.replace(old, new); k += 1
    assert k, f"slide {n}: no hay «{old}» en fórmulas"

# ---------------- notación: G en lugar de μ (13, 16, 19) ----------------
check(13, "En ese medio"); check(16, "Las velocidades de propagación"); check(19, "De la velocidad de corte")
for n in (13, 16, 19): math_replace(n, "𝜇", "𝐺")
plain_par(body_par(S(13), "Text 3", "Ecuación de movimiento")[1],
          [("Ecuación de movimiento. ", "n"), ("u", "i"), (": desplazamiento; ", "n"), ("ρ", "i"), (": densidad; ", "n"),
           ("λ", "i"), (", ", "n"), ("G", "i"), (": constantes de Lamé (", "n"), ("G", "i"), (" es el módulo de corte)", "n")])
plain_par(body_par(S(16), "Cuerpo", "𝐺: módulo")[1], [("G", "i"), (": módulo de corte, la rigidez al corte", "n")])

# ---------------- 35 (nueva): reflexión en la superficie libre ----------------
check(35, "Refracción: del gráfico")
set_title(S(35), "Reflexión de las ondas de cuerpo en la superficie libre")
old = shape(35, "Dromocrónica")
swap_picture(35, "Dromocrónica", os.path.join(FIG, "fig_reflexion_superficie_es.png"), "Reflexión superficie libre",
             box=(Emu(457200), old.top, Emu(8229600), Emu(2880000)))
tb = S(35).shapes.add_textbox(Emu(457200), old.top + Emu(2971800), Emu(8229600), Emu(457200))
tb.name = "Mensaje"; r = tb.text_frame.paragraphs[0].add_run()
r.text = "Una P o una SV generan P y SV reflejadas. Una SH genera solo SH."
r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79); r.font.name = "Arial"
set_cita(35, "Basada en Foti et al. (2014), cap. 2, p. 49, fig. 2.10. Ángulos según la ley de Snell con ν = 0,25.")
set_notes(S(35), "Antes de ver qué ondas se acoplan, veamos qué pasa cuando una onda de cuerpo llega a la superficie libre. "
                 "Por la ley de Snell, una P incidente genera una P reflejada con el mismo ángulo y una SV con un ángulo menor, "
                 "porque la S es más lenta. Una SV incidente genera, a su vez, una SV y una P. Es la conversión de modo. "
                 "La SH, en cambio, solo genera SH. Foti lo muestra en la figura 2.10 del capítulo 2.")

# ---------------- 36: conversión de modos, con fuente ----------------
check(36, "Conversión de modos")
set_cita(36, "Foti et al. (2014), cap. 2, p. 49.")

# ---------------- 39 (nueva): condición de existencia de Rayleigh (copia de la de Love) ----------------
check(39, "Condición de existencia de las ondas Love")
set_title(S(39), "Condición de existencia de las ondas Rayleigh")
swap_picture(39, "Figura condición Love", os.path.join(FIG, "fig_rayleigh_profundidad_es.png"), "Rayleigh con la profundidad")
remove(39, "Fondo thc"); remove(39, "Ecuación E_thc")
cue = unwrap(S(39), "Cuerpo")
cue.find('.//a:xfrm/a:off', NS).set('y', shape(39, "Rayleigh con la profundidad").top.__str__())
ps = cue.findall('.//a:p', NS); tpl = ps[0]
for p in ps[1:]: p.getparent().remove(p)
lines = ["Basta una superficie libre: existen incluso en un semiespacio homogéneo",
         "Surgen del acople de P y SV, que la superficie libre convierte entre sí",
         "Su amplitud se concentra en una profundidad del orden de una longitud de onda",
         "Ventaja frente a Love: no dependen de que haya capas"]
plain_par(tpl, lines[0])
for line in lines[1:]:
    q = copy.deepcopy(tpl); plain_par(q, line); cue.find('.//p:txBody', NS).append(q)
set_cita(39, "Foti et al. (2014), cap. 2, pp. 53–60 y fig. 2.16. Curva calculada con la solución de Lamb para ν = 0,25 (Achenbach, 1984).")
set_notes(S(39), "A diferencia de Love, las ondas Rayleigh no necesitan capas: basta la superficie libre. Existen incluso en un "
                 "semiespacio homogéneo, porque nacen del acople entre P y SV que la superficie libre impone. "
                 "La figura muestra cómo decae su amplitud con la profundidad: casi todo el movimiento queda dentro de "
                 "una longitud de onda, y la componente horizontal cambia de signo cerca de 0,2 longitudes de onda. "
                 "En superficie, la componente vertical es cerca de 1,5 veces la horizontal. Que existan siempre es una "
                 "ventaja práctica frente a las Love.")

# ---------------- 40: Rayleigh, movimiento (sin lo que pasó a la 39) ----------------
check(40, "Onda Rayleigh: origen")
set_title(S(40), "Onda Rayleigh: movimiento de las partículas")
for start in ("Acopla P y SV", "Existe incluso"):
    sp, p = body_par(S(40), "Cuerpo", start); p.getparent().remove(p)

# ---------------- 41 y 44: primer modelo y modelo estratificado ----------------
check(41, "Reparto de la energía")
set_title(S(41), "De vuelta al primer modelo: reparto de la energía de una fuente vertical")
check(44, "Perfil de velocidad de corte")
set_title(S(44), "Qué cambia en un modelo estratificado")

# ---------------- 56: tabla de transductores ----------------
check(56, "Transductores posibles")
tbl = next(sh for sh in S(56).shapes if sh.has_table).table
plain_par(tbl.cell(3, 3)._tc.find('.//a:p', NS), "Triaxial, salida digital. Ruido: 22,5 µg/√Hz")
plain_par(tbl.cell(4, 3)._tc.find('.//a:p', NS), "Banda muy amplia. Pesa 8,4 kg")
set_cita(56, "Sensor Nederland B.V. (2006); R. T. Clark (2026); Analog Devices (2026); Güralp Systems (s. f.); "
             "SparkFun Electronics (2026); DigiKey (2026).")

# ---------------- 61: respuesta de 10 Hz contra 4,5 Hz en lugar de los precios ----------------
check(61, "Por qué importa la pérdida")
x0 = shape(61, "Cabecera Geófono de 4,5 Hz").left; y0 = shape(61, "Cabecera Geófono de 4,5 Hz").top
x1 = shape(61, "Tarjeta SM-24 de 10 Hz").left + shape(61, "Tarjeta SM-24 de 10 Hz").width
for nm in ("Cabecera Geófono de 4,5 Hz", "Text 2", "Tarjeta Geófono de 4,5 Hz", "Text 4",
           "Cabecera SM-24 de 10 Hz", "Text 6", "Tarjeta SM-24 de 10 Hz", "Text 8"):
    remove(61, nm)
pic = S(61).shapes.add_picture(os.path.join(FIG, "fig_geofonos_10_vs_45_es.png"), x0, y0)
r = (x1 - x0) / pic.width; pic.width, pic.height = int(pic.width * r), int(pic.height * r); pic.name = "Geófonos 10 y 4,5 Hz"
shape(61, "Text 10").text_frame.paragraphs[0].runs[0].text = ""
set_cita(61, "Modelo de segundo orden con ζ = 0,6 para los dos geófonos (ilustrativo).")

# ---------------- 62: técnicas ----------------
check(62, "Existen cuatro familias")
shape(62, "Text 8").text_frame.paragraphs[0].runs[0].text = (
    "Se cancela la dinámica del sensor para reconstruir la baja frecuencia. Amplifica el ruido de baja frecuencia "
    "y lleva fácilmente a la inestabilidad.")
t16 = shape(62, "Text 16").text_frame.paragraphs[0].runs[0]
assert "se amplifica todo" in t16.text
t16.text = t16.text.replace("se amplifica todo", "se amplifica en su totalidad")

# ---------------- 75: offset ----------------
check(75, "El problema del offset")
plain_par(body_par(S(75), "Cuerpo", "Es consecuencia")[1],
          "Es consecuencia de las ganancias agresivas por etapa. Sin calibrar, el offset llegó hasta el 45,8 % del fondo de escala")
set_cita(75, "45,8 %: Álvarez, González y Vargas, DA-AFE (URUCON), fig. 4.")

# ---------------- 84: media móvil normal ----------------
check(84, "Diezmado por media móvil")
swap_picture(84, "Respuesta de la media móvil", os.path.join(FIG, "fig_media_movil_es.png"), "Respuesta de la media móvil")
set_cita(84, "Curva ilustrativa: media móvil de 4 muestras.")

# ---------------- 89: HAMMER, diagrama de bloques nuevo ----------------
check(89, "El nodo HAMMER")
swap_picture(89, "Principio ICP", os.path.join(FIG, "fig_hammer_bloques_es.png"), "Bloques del nodo HAMMER")

# ---------------- números de slide corridos (pies y notas) ----------------
for sl in prs.slides:
    for t in sl._element.iter(A + 't'):
        if t.text:
            t.text = t.text.replace("modelo de la slide 42", "modelo de la slide 44").replace("curva de la slide 44", "curva de la slide 46")
    if sl.has_notes_slide:
        tf = sl.notes_slide.notes_text_frame
        if "slide 58" in tf.text: tf.text = tf.text.replace("slide 58", "slide 60")

prs.save(deck)
print("ok")
