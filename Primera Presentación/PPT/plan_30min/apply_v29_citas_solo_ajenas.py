"""Quita de las citas toda autoría propia. Solo se citan fuentes ajenas.
Uso: python clean_cites.py --dry cites.json   |   python clean_cites.py in.pptx out.pptx"""
import sys, re, json

SELF = re.compile(r'propi[oa]s?\b|por el autor|aportad|\.py\b|\.tex\b|drawio|latex-|main\.c|resumen firmado|documento de tesis|del repo|no est[áa]n? en el resumen|no figura en el resumen|f_presupuesto|^Texto:?$', re.I)
AJENO = re.compile(r'\((19|20)\d\d[a-z]?\)|\b(19|20)\d\d\b|et al|Foti|Haeni|Ashlock|Luu|Nederland|PCB|Espressif|802\.11|Korsimaa|SparkFun|Clark|Breslin|Zhang|URUCON|PSoC Creator|Hakimitoroghi|Pazos|Wikimedia|EPA|SLB|Putt|Froldi|ASTM', re.I)

# Segmentos mixtos (propio + ajeno) o con matiz: reescritura a mano.
MIX = {
 "Imagen: esquema propio (fig_dromo_es.py), como en Haeni et al. (1987), fig. 2, p. 6": "Imagen: basada en Haeni et al. (1987), fig. 2, p. 6",
 "Esquema propio sobre el procedimiento general de inversión (Foti et al., 2014, cap. 1, fig. 1.4, p. 9, y cap. 6)": "Procedimiento general de inversión: Foti et al. (2014), cap. 1, fig. 1.4, p. 9, y cap. 6",
 "Cálculo propio con disba (Luu, 2021), modelo de la slide 39": "Calculado con disba (Luu, 2021), modelo de la slide 39",
 "Modelo propio a partir de Sensor Nederland B.V. (2006), circuito abierto": "Modelo a partir de Sensor Nederland B.V. (2006), circuito abierto",
 "Elaboración propia en PSoC Creator (resumen firmado, fig. 5 y sec. 4.3)": "Diseño en PSoC Creator",
 "Elaboración propia en PSoC Creator (resumen firmado, fig. 10 y sec. 4.6)": "Diseño en PSoC Creator",
 "Izquierda: esquema propio con datos de 03_excitacion.tex": None,
 "Ilustración propia (esquemática)": None,
 "Álvarez, González y Vargas, DA-AFE (URUCON), fig. 2, traducida desde su archivo fuente (drawio)": "Álvarez, González y Vargas, DA-AFE (URUCON), fig. 2, traducida",
 "Documento de tesis: 04_arquitectura.tex, 08_trabajos_futuros.tex (UWB, Korsimaa 2025), 01_problema_modelado.tex (GPS)": "UWB: Korsimaa et al. (2025)",
 "Wi-Fi, LoRa y GPS: propiedades generales, no medidas propias": "Wi-Fi, LoRa y GPS: propiedades generales, no medidas",
 "Documento de tesis, 04_arquitectura.tex (sec. broadcast vs unicast)": None,
 "comportamiento de ACK en unicast: estándar 802.11": "Comportamiento de ACK en unicast: estándar IEEE 802.11",
 "documento de tesis, 04_validacion.tex (Lin y Ashlock, 2016)": "Lin y Ashlock (2016)",
 "Documento de tesis, 02_requerimientos.tex y 04_validacion.tex (atribuye a Foti et al., 2018, pp. 2380 y 2382: L ≳ λmáx y L ≳ 1,5·λmáx": "Foti et al. (2018), pp. 2380 y 2382: L ≳ λmáx y L ≳ 1,5·λmáx",
 "L ≳ 2·λmáx es criterio de la tesis)": "L ≳ 2·λmáx es criterio adoptado, no de Foti",
 "Figura propia (fig/fig_cuna.py), con c = 150 m/s": "Con c = 150 m/s",
 "Figura propia ilustrativa (fig/fig_delta_sigma_es.py): modulador de 2.º orden y filtro sinc³ genéricos, no los del PSoC": "Figura ilustrativa: modulador de 2.º orden y filtro sinc³ genéricos, no los del PSoC",
 "Registro sintético propio (ilustrativo), con la curva de la slide 41": "Registro sintético (ilustrativo), con la curva de la slide 41",
 "Los 13,8 dB son 20·log10(5/1,024), cálculo propio": "Los 13,8 dB son 20·log10(5/1,024)",
 "Curva ilustrativa: cálculo propio de la respuesta de una media móvil": "Curva ilustrativa: respuesta de una media móvil",
 "estabilidad y ruido: inferencia propia, por confirmar": "estabilidad y ruido: por confirmar",
 "Esquema genérico aportado por el autor (fuente por confirmar)": "Esquema genérico (fuente por confirmar)",
 "ganancia: documento de tesis, sec. 3": None,
 "Precios de referencia consultados por el autor en 2026": "Precios de referencia de 2026",
 "El límite de 512 lotes cuenta muestras ya diezmadas (firmware del PSoC, main.c)": "El límite de 512 lotes cuenta muestras ya diezmadas",
 "derecha: PSoC Creator": "Derecha: PSoC Creator",
}
FULL = {
 "Resumen firmado, ec. 16 y 17. Reloj: AcondicionamientoAnalogico.cydwr (ADC:theACLK ← MASTER_CLK ← PLL ← IMO; XTAL deshabilitado; ±1 %). 10 y 50 ppm: ilustrativos.":
 "Reloj del ADC: IMO → PLL → MASTER_CLK, sin cristal; PSoC Creator le asigna ±1 %. 10 y 50 ppm: ilustrativos.",
}

# Cambios en el cuerpo de las slides, no en las citas.
BODY = [(" (propuesta propia)", ""), (", misma geometría (resumen)", ", misma geometría"), ("INFERENCIA PROPIA", "INFERENCIA")]

SPLIT = re.compile(r'(; |\. (?=[A-ZÁÉÍÓÚÑ¿«])|\.$| \| )')

def strong(a, b):
    return ". " if (a.startswith(".") or b.startswith(".") or "|" in a or "|" in b) else a

def clean(text):
    for a, b in BODY:
        if a in text: return text.replace(a, b), []
    pre = ""
    if text.startswith("Fuentes: ") and text[9:] in FULL: return "Fuentes: " + FULL[text[9:]], []
    if text in FULL: return FULL[text], []
    if text.startswith("Fuentes: "): pre, text = "Fuentes: ", text[9:]
    parts = SPLIT.split(text)
    segs = parts[0::2]; seps = parts[1::2] + [""]
    keep, todo = [], []
    for seg, sep in zip(segs, seps):
        s = seg.strip()
        if s in MIX:
            new = MIX[s]
            if new: keep.append([new, sep])
            elif keep: keep[-1][1] = strong(keep[-1][1], sep)
            continue
        if SELF.search(s):
            if AJENO.search(s): todo.append(s); keep.append([seg, sep])
            if keep: keep[-1][1] = strong(keep[-1][1], sep)
            continue  # segmento solo propio: se quita
        if s or sep.strip() == "|": keep.append([seg, sep])
    # rearmar: separador final de cada segmento; cerrar con punto
    out = ""
    for k, (seg, sep) in enumerate(keep):
        last = k == len(keep) - 1
        if last: sep = "" if seg.rstrip().endswith(".") else "."
        elif sep.strip() == "": sep = ". "
        if not last and seg.rstrip().endswith(".") and sep.startswith("."): sep = " "
        if k and out.endswith(". "): seg = seg.strip(); seg = seg[:1].upper() + seg[1:]
        out += seg.strip() if k == 0 else (" " + seg.strip() if not out.endswith(" ") else seg.strip())
        out += sep.rstrip() if last else sep
    out = re.sub(r'\s+\|\s*$', '', out).strip()
    out = re.sub(r'^\|\s*', '', out)
    out = out[:1].upper() + out[1:] if out else out
    if out in (".", ""): return "", todo
    return pre + out, todo

if __name__ == "__main__" and sys.argv[1] == "--dry":
    d = json.load(open(sys.argv[2], encoding="utf-8"))
    for e in d:
        if e["math"]: print("MATH   ", e["where"][:2], e["text"]); continue
        new, todo = clean(e["text"])
        print("OLD", e["where"][:2], e["text"]); print("NEW", new if new else "<<VACÍO>>")
        for t in todo: print("  TODO:", t)
        print()
elif __name__ == "__main__":
    import pptx
    prs = pptx.Presentation(sys.argv[1]); n = vac = 0; todos = []
    for i, s in enumerate(prs.slides, 1):
        stack = list(s.shapes)
        while stack:
            sh = stack.pop()
            if sh.shape_type == 6: stack.extend(sh.shapes); continue
            if not sh.has_text_frame: continue
            for p in sh.text_frame.paragraphs:
                if 'oMath' in p._p.xml or 'AlternateContent' in p._p.xml or len(p.runs) != 1: continue
                t = p.runs[0].text
                if not (SELF.search(t) or any(a in t for a, _ in BODY)) or re.match(r'^(Mismas propiedades|offset propio|El maestro crea su propio|Cada nodo acondiciona|Ondas: usan propiedades)', t.strip()): continue
                new, todo = clean(t); todos += [(i, x) for x in todo]
                if new != t: p.runs[0].text = new; n += 1; vac += (new == "")
    prs.save(sys.argv[2]); print("cambiados:", n, "vaciados:", vac, "TODO:", todos)

# Párrafos con fórmulas: se editan por run o se vacían enteros.
MATH_SUB = [("Modelo de ejemplo propio (ilustrativo)", "Modelo de ejemplo (ilustrativo)"),
            ("documento de tesis, ec. (comp-afe); ", ""),
            ("Documento de tesis, 04_arquitectura.tex (criterio de reparto); resumen firmado, sec. 4.6. ", "")]
MATH_BLANK = ("Cálculo propio con 512 lotes", "Figura propia (f_presupuesto_sincronizacion)")

def fix_math(prs):
    from pptx.oxml.ns import qn
    n = 0
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame: continue
            for p in sh.text_frame.paragraphs:
                full = "".join(r.text for r in p.runs)
                if any(full.lstrip("Fuentes: ").startswith(b) or b in full for b in MATH_BLANK):
                    keep_tags = (qn("a:pPr"), qn("a:endParaRPr"))
                    for ch in list(p._p):
                        if ch.tag not in keep_tags: p._p.remove(ch)
                    if full.startswith("Fuentes: "): p.add_run().text = ""
                    n += 1; continue
                for r in p.runs:
                    for a, b in MATH_SUB:
                        if a in r.text: r.text = r.text.replace(a, b); n += 1
    return n

if __name__ == "__main__" and len(sys.argv) == 3 and sys.argv[1] != "--dry":
    import pptx
    prs = pptx.Presentation(sys.argv[2]); print("fórmulas:", fix_math(prs)); prs.save(sys.argv[2])
