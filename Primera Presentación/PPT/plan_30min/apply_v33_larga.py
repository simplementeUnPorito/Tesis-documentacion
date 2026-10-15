"""v33 larga: la 64 se separa en «Resistencia negativa» (64) y «Balance de fuerza» (65, copia de la 64)."""
import sys, os, copy
import pptx
from pptx_helpers_v30 import NS, A, ptext, top_shapes, unwrap, plain_par, body_par, set_title, set_notes

deck = sys.argv[1]
FIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figuras_es_v25")
prs = pptx.Presentation(deck)
S = lambda n: prs.slides[n - 1]
shape = lambda n, nm: next(sh for sh in S(n).shapes if sh.name == nm)
title = lambda n: ptext(top_shapes(S(n), "Título")[1][0].find('.//a:p', NS))
for n in (64, 65): assert title(n).startswith("Resistencia negativa y balance"), (n, title(n))

def set_text(n, name, text):
    p = shape(n, name).text_frame.paragraphs[0]
    if p.runs: p.runs[0].text = text; [r._r.getparent().remove(r._r) for r in p.runs[1:]]
    else: p.add_run().text = text

def bullets(n, lines):
    sp = unwrap(S(n), "Cuerpo"); tx = sp.find('.//p:txBody', NS); ps = tx.findall('a:p', NS)
    for p in ps[1:]: tx.remove(p)
    plain_par(ps[0], lines[0])
    for line in lines[1:]:
        q = copy.deepcopy(ps[0]); plain_par(q, line); tx.append(q)
    return sp

# ---- 64: resistencia negativa (conserva la ecuación del amortiguamiento) ----
set_title(S(64), "Resistencia negativa")
bullets(64, [[("Una carga activa ", "n"), ("Z", "i"), ("L", "s"), (" negativa cancela parte de ", "n"), ("R", "i"), ("i", "s"),
              (" y sube el amortiguamiento", "n")],
             [("Es un lazo con realimentación positiva: si ", "n"), ("R", "i"), ("i", "s"),
              (" cambia con la temperatura, puede volverse inestable", "n")],
             "A frecuencias altas, la inductancia de la bobina distorsiona la respuesta"])
set_text(64, "Text 4", "")
set_text(64, "Cita", "Ma et al. (2023); Zhang et al. (2025), Chin. J. Electron. 34(6), 1678–1688.")
set_notes(S(64), "La primera de las dos técnicas activas es la resistencia negativa. Un convertidor de impedancia negativa "
                 "presenta al geófono una carga que cancela parte de la resistencia de la bobina, y así sube el amortiguamiento "
                 "electromecánico, como muestra la ecuación. El problema es que ese convertidor es un lazo con realimentación "
                 "positiva: si la resistencia de la bobina cambia con la temperatura, el conjunto puede volverse inestable. "
                 "Además, Ma señala que a frecuencias altas la inductancia de la bobina ya no se puede ignorar y distorsiona la respuesta.")

# ---- 65: balance de fuerza (figura del lazo en lugar de la ecuación) ----
set_title(S(65), "Balance de fuerza")
for nm in ("Fondo E_zeta2", "Ecuación E_zeta2"):
    cont, _ = top_shapes(S(65), nm); cont.getparent().remove(cont)
cue = unwrap(S(65), "Cuerpo")
tit, msg = shape(65, "Título"), shape(65, "Text 4")
top = tit.top + tit.height
area_h = msg.top - top
pic = S(65).shapes.add_picture(os.path.join(FIG, "fig_balance_fuerza_es.png"), shape(65, "Cita").left, top)
r = (area_h * 0.95) / pic.height; pic.width, pic.height = int(pic.width * r), int(pic.height * r); pic.name = "Lazo de balance de fuerza"
off = cue.find('.//a:xfrm/a:off', NS); ext = cue.find('.//a:xfrm/a:ext', NS)
x0 = pic.left + pic.width + 182880
off.set('x', str(x0)); off.set('y', str(top + 91440)); ext.set('cx', str(8686800 - x0)); ext.set('cy', str(area_h - 91440))
bullets(65, ["Un lazo mide la masa e inyecta corriente en la bobina",
             "La fuerza de Lorentz extiende la banda hacia bajas frecuencias",
             "El lazo interactúa con el transductor: hay que resolver estabilidad, ruido y tolerancias entre unidades"])
set_text(65, "Cita", "Zhang et al. (2025), Chin. J. Electron. 34(6), 1678–1688; Song et al. (2016), según Ma et al. (2023).")
set_notes(S(65), "La segunda técnica activa es el balance de fuerza. Un lazo mide el movimiento de la masa y calcula una corriente "
                 "que se inyecta en la misma bobina. La fuerza de Lorentz se opone al movimiento y la banda se extiende hacia bajas "
                 "frecuencias, como muestran Zhang y Song. Funciona, pero el lazo queda pegado al transductor: hay que resolver "
                 "estabilidad, ganancia, ruido, tolerancias y que todas las unidades se comporten igual. Para un nodo multicanal "
                 "barato y replicable, ese costo de complejidad pesa.")

prs.save(deck); print("ok")
