"""Comprimida v37 a partir de la larga v37, sin portapapeles.
Cada slide comprimida reúne varias slides de la larga como escenas, una por clic:
la escena 0 está visible al entrar; en cada clic salen (fundido) los objetos de la escena anterior y entran
(fundido) los de la siguiente. Las relaciones de cada objeto (imágenes, gráficos, respaldos de ecuaciones,
hipervínculos) se reapuntan desde la slide base a las mismas partes, y los ids se renumeran.
Las notas concatenan las de las escenas con el rótulo «[Slide larga NN: título]».
Uso: python build_comprimida_v37.py larga.pptx salida.pptx"""
import sys, copy, re
import pptx
from lxml import etree
from pptx_helpers_v30 import NS, P, A, ptext, top_shapes

GROUPS = [
    [1], [2, 3, 4], [5, 6, 7, 8, 9, 10], [11, 12, 13], [14, 15, 16, 17, 18, 19, 20], [21, 22, 23], [24, 25, 26, 27],
    [28, 29, 30, 31, 32], [33, 34, 35, 36], [37, 38, 39, 40, 41, 42, 43], [44, 45, 46, 47, 48], [49, 50, 51, 52, 53],
    [54, 55, 56], [57, 58, 59, 60], [61, 62, 63, 64, 65], [66, 67, 68], [69, 70, 71, 72, 73, 74], [75, 76, 77, 78, 79],
    [80, 81], [82, 83, 84, 85, 86], [87, 88, 89, 90, 91], [92, 93, 94, 95, 96], [97, 98, 99], [100, 101, 102, 103, 104],
    [105, 106, 107], [108, 109, 110], [111, 112, 113], [114, 115, 116, 117], [118], [119], [120, 121, 122, 123, 124],
    [125, 126],
]
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
MC = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
EXIT_MS, ENTER_MS, ENTER_DELAY = 250, 400, 250

src, dst = sys.argv[1], sys.argv[2]
prs = pptx.Presentation(src)
slides = list(prs.slides)
flat = [n for g in GROUPS for n in g]
assert sorted(flat) == list(range(1, len(slides) + 1)), "los grupos no cubren la larga exactamente una vez"

def is_number_placeholder(el):
    return el.tag == P + 'sp' and el.find('.//p:nvPr/p:ph', NS) is not None

def top_children(slide):
    tree = slide._element.find('.//p:spTree', NS)
    return [ch for ch in tree if ch.tag not in (P + 'nvGrpSpPr', P + 'grpSpPr')]

def top_id(el):
    """id del objeto de nivel superior (en AlternateContent, el de la copia Choice)."""
    c = next(el.iter(P + 'cNvPr'))
    return int(c.get('id'))

def title_of(slide):
    try: return ptext(top_shapes(slide, "Título")[1][0].find('.//a:p', NS))
    except Exception:
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip(): return sh.text_frame.text.strip().split("\n")[0]
    return ""

def notes_of(slide):
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""

# ---------------- animación (fragmentos de anim2.py) ----------------
def build_timing(scenes):
    n = [2]
    def nid(): n[0] += 1; return n[0]
    tgt = lambda s: f'<p:tgtEl><p:spTgt spid="{s}"/></p:tgtEl>'
    def set_vis(s, val, delay=0):
        return (f'<p:set><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst></p:cTn>'
                f'{tgt(s)}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:to><p:strVal val="{val}"/></p:to></p:set>')
    def wrap(cls, delay, first, inner):
        nt = 'clickEffect' if first else 'withEffect'
        return (f'<p:par><p:cTn id="{nid()}" presetID="10" presetClass="{cls}" presetSubtype="0" fill="hold" nodeType="{nt}">'
                f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>')
    def fx_exit(s, first):
        inner = (f'<p:animEffect transition="out" filter="fade"><p:cBhvr><p:cTn id="{nid()}" dur="{EXIT_MS}"/>{tgt(s)}</p:cBhvr></p:animEffect>'
                 + set_vis(s, 'hidden', EXIT_MS - 1))
        return wrap('exit', 0, first, inner)
    def fx_enter(s, first):
        inner = set_vis(s, 'visible') + \
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{nid()}" dur="{ENTER_MS}"/>{tgt(s)}</p:cBhvr></p:animEffect>'
        return wrap('entr', ENTER_DELAY, first, inner)
    clicks = []
    for k in range(1, len(scenes)):
        effs = []; first = True
        for s in scenes[k - 1]: effs.append(fx_exit(s, first)); first = False
        for s in scenes[k]: effs.append(fx_enter(s, first)); first = False
        o1, o2 = nid(), nid()
        clicks.append(f'<p:par><p:cTn id="{o1}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                      f'<p:par><p:cTn id="{o2}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{"".join(effs)}</p:childTnLst></p:cTn></p:par>'
                      '</p:childTnLst></p:cTn></p:par>')
    xml = ('<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
           '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>' + ''.join(clicks) +
           '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    cnt = [0]
    def ren(m): cnt[0] += 1; return f'<p:cTn id="{cnt[0]}"'
    return etree.fromstring(re.sub(r'<p:cTn id="\d+"', ren, xml))

# ---------------- fusión ----------------
absorbed = []
log = []
for g in GROUPS:
    base = slides[g[0] - 1]
    btree = base._element.find('.//p:spTree', NS)
    ids_in_base = [int(c.get('id')) for c in base._element.iter(P + 'cNvPr') if c.get('id', '').isdigit()]
    next_id = max(ids_in_base + [1]) + 1
    scenes = [[top_id(el) for el in top_children(base) if not is_number_placeholder(el)]]
    notes = [f"[Slide larga {g[0]}: {title_of(base)}] {notes_of(base)}".strip()]
    num_ph = next((el for el in top_children(base) if is_number_placeholder(el)), None)
    for k, m in enumerate(g[1:], 1):
        ms = slides[m - 1]
        notes.append(f"[Slide larga {m}: {title_of(ms)}] {notes_of(ms)}".strip())
        scene = []
        for el in top_children(ms):
            if is_number_placeholder(el): continue
            new = copy.deepcopy(el)
            # relaciones: misma parte de destino, nueva relación desde la base
            for node in new.iter():
                for attr in list(node.attrib):
                    if attr.startswith('{%s}' % R):
                        old = node.get(attr); rel = ms.part.rels[old]
                        if rel.is_external: rid = base.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
                        else: rid = base.part.relate_to(rel.target_part, rel.reltype)
                        node.set(attr, rid)
            # ids: Choice y Fallback comparten id; los conectores apuntan por id
            remap = {}
            for c in new.iter(P + 'cNvPr'):
                o = c.get('id')
                if o not in remap: remap[o] = str(next_id); next_id += 1
                c.set('id', remap[o]); c.set('name', f"E{k} " + c.get('name', ''))
            for tag in (A + 'stCxn', A + 'endCxn'):
                for c in new.iter(tag):
                    if c.get('id') in remap: c.set('id', remap[c.get('id')])
            if num_ph is not None: num_ph.addprevious(new)
            else: btree.append(new)
            scene.append(top_id(new))
        scenes.append(scene)
        absorbed.append(ms)
    if len(scenes) > 1:
        old_t = base._element.find('p:timing', NS)
        if old_t is not None: base._element.remove(old_t)
        t = build_timing(scenes)
        ext = base._element.find('p:extLst', NS)
        (ext.addprevious(t) if ext is not None else base._element.append(t))
    if base.has_notes_slide or len(g) > 1:
        base.notes_slide.notes_text_frame.text = "\n\n".join(notes)
    log.append(f"larga {','.join(map(str, g))} ({len(g)} escenas)")

# borrar las slides absorbidas
sld_lst = prs.slides._sldIdLst
by_part = {s.part: s for s in absorbed}
for sid in list(sld_lst):
    rid = sid.get('{%s}id' % R)
    if prs.part.related_part(rid) in by_part:
        prs.part.drop_rel(rid); sld_lst.remove(sid)

prs.save(dst)
for i, l in enumerate(log, 1): print(f"{i:2d}: {l}")
print("total:", len(prs.slides), "slides")
