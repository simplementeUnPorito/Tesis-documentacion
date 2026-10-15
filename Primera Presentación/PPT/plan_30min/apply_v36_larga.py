"""v36 larga: la 64 usa la fig. 2 de X. Zhang et al. (2020), traducida por Elías, en lugar del esquema propio.
Las dos citas «Zhang» se distinguen por la inicial (X. Zhang 2020, Z. Zhang 2025)."""
import sys, os, copy, re
import pptx
from pptx.util import Pt
from pptx_helpers_v30 import NS, A, ptext, top_shapes, unwrap, plain_par, set_notes

deck = sys.argv[1]
FIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figuras_es_v25")
prs = pptx.Presentation(deck)
s = prs.slides[63]
assert ptext(top_shapes(s, "Título")[1][0].find('.//a:p', NS)).startswith("Balance de fuerza")
sh = {x.name: x for x in s.shapes}
tit, msg, cita = sh["Título"], sh["Text 4"], sh["Cita"]

# figura: tan ancha como permita la altura libre
old = sh["Lazo de balance de fuerza"]
top = tit.top + tit.height + Pt(4); bottom = Pt(296)
pic = s.shapes.add_picture(os.path.join(FIG, "zhang2020_fig2_traducida.png"), Pt(36), top)
r = min(Pt(648) / pic.width, (bottom - top) / pic.height)
pic.width, pic.height = int(pic.width * r), int(pic.height * r)
pic.left = Pt(36) + (Pt(648) - pic.width) // 2; pic.top = top
pic.name = "Lazo de balance de fuerza"
old._element.addprevious(pic._element); old._element.getparent().remove(old._element)

# una línea con el dato del artículo y el mensaje en una línea
cue = unwrap(s, "Cuerpo"); tx = cue.find('.//p:txBody', NS); ps = tx.findall('a:p', NS)
for p in ps[1:]: tx.remove(p)
plain_par(ps[0], "Un lazo PID por la misma bobina lleva un geófono de 4,5 Hz hasta 0,16 Hz")
pPr = ps[0].find('a:pPr', NS)
if pPr is not None:
    for k in ("marL", "indent"): pPr.attrib.pop(k, None)
    bu = pPr.find('a:buChar', NS)
    if bu is not None:                     # una sola línea de texto: sin viñeta
        from lxml import etree
        bu.addprevious(etree.Element(A + 'buNone')); pPr.remove(bu)
for rp in cue.iter(A + 'rPr'): rp.set('sz', '1600')
off = cue.find('.//a:xfrm/a:off', NS); ext = cue.find('.//a:xfrm/a:ext', NS)
off.set('x', str(Pt(36))); off.set('y', str(Pt(300))); ext.set('cx', str(Pt(648))); ext.set('cy', str(Pt(22)))
msg.top, msg.left, msg.width, msg.height = Pt(324), Pt(36), Pt(648), Pt(26)
mp = msg.text_frame.paragraphs[0]
mp.runs[0].text = "Funciona, pero su complejidad pesa en un nodo multicanal barato y replicable."
for r_ in mp.runs[1:]: r_._r.getparent().remove(r_._r)
cp = cita.text_frame.paragraphs[0]
cp.runs[0].text = "Imagen: X. Zhang et al. (2020), Sens. Actuators A 312, 112074, fig. 2, traducida. Ver también Z. Zhang et al. (2025)."
for r_ in cp.runs[1:]: r_._r.getparent().remove(r_._r)
set_notes(s, "La segunda técnica activa es el balance de fuerza. La figura es el modelo de Zhang y colaboradores, de 2020: "
             "la misma bobina del geófono mide y actúa. La tensión de la bobina pasa por un modulador sigma-delta de cuarto "
             "orden; el tren de bits vuelve por un DAC de un bit y un controlador PID, y la bobina aplica una fuerza de Lorentz "
             "que se opone al movimiento. Con eso llevan un geófono comercial de 4,5 hertz hasta 0,16 hertz. Funciona, pero "
             "es un lazo pegado al transductor, con analógico y digital, y para un nodo multicanal barato y replicable ese "
             "costo de complejidad pesa.")

# las dos citas «Zhang» se distinguen por la inicial (APA)
for sl in prs.slides:
    for t in sl._element.iter(A + 't'):
        if t.text and "Zhang et al. (2025)" in t.text and "Z. Zhang" not in t.text:
            t.text = t.text.replace("Zhang et al. (2025)", "Z. Zhang et al. (2025)")

# referencia nueva
boxes = [x for sl in prs.slides for x in sl.shapes if x.name == "Referencias"]
tpl = boxes[-1].text_frame.paragraphs[0]._p
e = copy.deepcopy(tpl); rs = e.findall(A + 'r')
rs[0].find(A + 't').text = ("Zhang, X., Wei, X. et al. (2020). A digital low-frequency geophone based on 4th-order sigma-delta "
                            "modulator and single-coil velocity feedback. Sensors and Actuators A: Physical, 312, 112074.")
for r_ in rs[1:]: e.remove(r_)
boxes[-1].text_frame._txBody.append(e)

prs.save(deck); print("ok")
