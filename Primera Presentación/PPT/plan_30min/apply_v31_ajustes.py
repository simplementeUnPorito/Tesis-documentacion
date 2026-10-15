"""v31 larga: ajustes visuales de 13, 39 y 89."""
import sys, pptx
from pptx.util import Pt
from pptx_helpers_v30 import NS, ptext, top_shapes, unwrap, plain_par, body_par

deck = sys.argv[1]
prs = pptx.Presentation(deck)
S = lambda n: prs.slides[n - 1]
shape = lambda n, nm: next(sh for sh in S(n).shapes if sh.name == nm)

plain_par(body_par(S(13), "Text 3", "Ecuación de movimiento")[1],
          [("Ecuación de movimiento. ", "n"), ("u", "i"), (": desplazamiento; ", "n"), ("ρ", "i"), (": densidad; ", "n"),
           ("λ", "i"), (", ", "n"), ("G", "i"), (": constantes de Lamé", "n")])

cue = unwrap(S(39), "Cuerpo")
for old, new in (("Basta una superficie libre: existen incluso en un semiespacio homogéneo", "Basta una superficie libre, incluso en un semiespacio homogéneo"),
                 ("Surgen del acople de P y SV, que la superficie libre convierte entre sí", "Surgen del acople de P y SV en la superficie libre"),
                 ("Su amplitud se concentra en una profundidad del orden de una longitud de onda", "Su amplitud se concentra en una longitud de onda de profundidad"),
                 ("Ventaja frente a Love: no dependen de que haya capas", "Ventaja frente a Love: no requieren capas")):
    plain_par(body_par(S(39), "Cuerpo", old)[1], new)
shape(39, "Cuerpo").height = shape(39, "Cita").top - Pt(6) - shape(39, "Cuerpo").top   # hasta el pie

blk = shape(89, "Bloques del nodo HAMMER"); sch = shape(89, "Acondicionamiento del martillo")
tit = shape(89, "Título"); cue = shape(89, "Cuerpo")
top = tit.top + tit.height + Pt(4); h = cue.top - Pt(6) - top          # zona libre entre título y viñetas
w = int(h * blk.width / blk.height)
blk.left, blk.top, blk.width, blk.height = Pt(36), top, w, h
r = sch.height / sch.width
sch.left = Pt(36) + w + Pt(16); sch.width = Pt(684) - sch.left; sch.height = int(sch.width * r)
if sch.height > h: sch.height = h; sch.width = int(h / r)
sch.top = top + (h - sch.height) // 2

prs.save(deck)
print("ok")
