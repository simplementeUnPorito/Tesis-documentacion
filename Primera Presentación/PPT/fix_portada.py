import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

IMG="masw-geostru-ES.png"
REF="Fuente: GEOSTRU, «MASW – Multichannel Analysis of Surface Waves», 2019. https://www.geostru.eu/es/blog/2019/06/02/masw/"
PANEL_W=Inches(3.3)

def txt(sh,text,size,color,bold=False,font=None,align=PP_ALIGN.LEFT,italic=False):
    tf=sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    for m in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(tf,m,0)
    p=tf.paragraphs[0]; p.alignment=align
    for r in list(p.runs): r._r.getparent().remove(r._r)
    r=p.add_run(); r.text=text; f=r.font; f.size=Pt(size); f.bold=bold; f.italic=italic
    f.color.rgb=RGBColor.from_string(color)
    if font: f.name=font

def fix(path):
    prs=Presentation(path); s=prs.slides[0]
    t,a,tu,pic,ph=list(s.shapes)
    s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor(255,255,255)
    # panel izquierdo
    panel=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,PANEL_W,prs.slide_height)
    panel.fill.solid(); panel.fill.fore_color.rgb=RGBColor(0x1E,0x27,0x61); panel.line.fill.background()
    s.shapes._spTree.remove(panel._element); s.shapes._spTree.insert(2,panel._element)
    L=Inches(0.3); W=PANEL_W-Inches(0.6)
    for sh,(y,h) in ((t,(0.55,2.7)),(a,(3.45,0.75)),(tu,(4.3,0.5))): 
        sh.left,sh.top,sh.width,sh.height=L,Inches(y),W,Inches(h)
    txt(t,"Diseño e implementación de un sistema electrónico multicanal del tipo IoT para caracterización geotécnica del suelo mediante ondas sísmicas",17,"FFFFFF",True,"Cambria")
    txt(a,"Elías David Álvarez Martínez | Ingeniería Electrónica | Y24127",11,"F5F5F5")
    txt(tu,"Tutor: Enrique A. Vargas Cabral, Ph.D.",11,"9AD7E0")
    # figura: todo el ancho disponible a la derecha, relación 2:1
    sp=s.shapes._spTree
    for sh in (pic,ph): sp.remove(sh._element)
    m=Inches(0.2); x=PANEL_W+m; w=prs.slide_width-x-m; h=int(w*1274/2548)
    top=int((prs.slide_height-h)/2)-Inches(0.2)
    s.shapes.add_picture(IMG,x,top,w,h)
    r=s.shapes.add_textbox(x,top+h+Inches(0.08),w,Inches(0.45))
    txt(r,REF,8,"595959",italic=True)
    prs.save(path)
for p in sys.argv[1:]: fix(p)
